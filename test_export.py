"""Export categories and row selection without touching the real scan database."""
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import marketplace_prescreener as app
import review_categories


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        db = Path(self.folder.name) / "test.sqlite3"
        self.db_patch = patch.object(app, "DB_PATH", db)
        self.db_patch.start()
        self.addCleanup(self.db_patch.stop)
        with sqlite3.connect(db) as conn:
            conn.execute("create table scan_history (job_id text primary key, created_at text, finished_at text, status text, criteria_json text)")
            conn.execute("create table reviews (id integer primary key, job_id text, scan_status text, "
                         "review_state text, country_signal text, keyword text, title text, listing_url text, negatives text not null default '', review_category text not null default '', created_at text not null default '')")
            conn.executemany("insert into reviews (id,job_id,scan_status,review_state,country_signal,keyword,title,listing_url,negatives) values (?, ?, ?, ?, ?, ?, ?, ?, ?)", [
                (1, "scan-a", "match", "pending", "China", "pending", "item-1", "url-1", "tool, refill"),
                (2, "scan-a", "match", "approved", "China", "approved", "item-2", "url-2", "tool, refill"),
                (3, "scan-a", "match", "rejected", "China", "disputed", "item-3", "url-3", "tool, refill"),
                (4, "scan-a", "unknown", "not_queued", "China", "not queued", "item-4", "url-4", "tool, refill"),
                (5, "scan-b", "match", "approved", "China", "other scan", "item-5", "url-5", ""),
            ])

    def export_keywords(self, job_id=None, scope="all", selected_ids=None, category=""):
        path = app.export_reviews_xlsx(job_id=job_id, scope=scope, selected_ids=selected_ids, category=category)
        self.addCleanup(path.unlink, missing_ok=True)
        workbook = load_workbook(path, read_only=True)
        try:
            rows = list(workbook.active.values)
        finally:
            workbook.close()
        self.assertIn("Review Status", rows[0])
        self.assertIn("Review Category", rows[0])
        self.assertEqual(rows[0][:2], ("Keyword", "Negatives"))
        return rows[1:]

    def test_scopes_and_scan_boundary(self):
        for scope, expected in {
            "all": {"pending", "approved", "disputed", "not queued"},
            "match": {"pending", "approved", "disputed"},
            "approved": {"approved"}, "pending": {"pending"},
            "rejected": {"disputed"}, "not_queued": {"not queued"},
        }.items():
            with self.subTest(scope=scope):
                self.assertEqual({row[0] for row in self.export_keywords("scan-a", scope)}, expected)
        self.assertEqual(len(self.export_keywords()), 5)

    def test_negatives_in_scoped_and_selected_exports(self):
        self.assertEqual(self.export_keywords("scan-a", "approved", {2})[0][1], "tool, refill")
        self.assertIsNone(self.export_keywords("scan-b", "approved")[0][1])

    def test_selected_rows_intersect_scope_and_job(self):
        self.assertEqual([row[0] for row in self.export_keywords("scan-a", "approved", {1, 2, 5})], ["approved"])
        self.assertEqual([row[0] for row in self.export_keywords("scan-a", "match", {3, 4, 5})], ["disputed"])
        self.assertEqual(self.export_keywords("scan-a", "all", set()), [])

    def test_category_filter_intersects_status_scan_and_selection(self):
        first = 'Trademark Infringement Media'
        second = 'IP Rights Infringement Utility Outdoor & Maintenence'
        with sqlite3.connect(app.DB_PATH) as conn:
            conn.executemany('update reviews set review_category = ? where id = ?',
                             [(first, 1), (first, 2), (second, 4), (first, 5)])
        self.assertEqual([row[0] for row in self.export_keywords('scan-a', 'approved', category=first)],
                         ['approved'])
        self.assertEqual(self.export_keywords('scan-a', 'approved', category=second), [])
        self.assertEqual({row[0] for row in self.export_keywords('scan-a', 'all', category=first)},
                         {'pending', 'approved'})
        self.assertEqual([row[0] for row in self.export_keywords('scan-a', 'not_queued', category=second)],
                         ['not queued'])
        self.assertEqual(self.export_keywords('scan-a', 'approved', {2}, category=second), [])
        self.assertEqual([row[0] for row in self.export_keywords('scan-a', 'approved', {2, 5}, first)],
                         ['approved'])
        with self.assertRaises(ValueError):
            self.export_keywords('scan-a', 'approved', category='Not a real category')

    def test_category_filter_via_download_routes(self):
        category = review_categories.CATEGORIES[0]
        with sqlite3.connect(app.DB_PATH) as conn:
            conn.execute('update reviews set review_category = ? where id in (2, 5)', (category,))
        response = app.export_xlsx(job='scan-a', scope='approved', category=category)
        self.addCleanup(response.path.unlink, missing_ok=True)
        sheet = list(load_workbook(response.path, read_only=True).active.values)
        self.assertEqual(len(sheet), 2)
        self.assertEqual(sheet[1][sheet[0].index('Review Category')], category)
        response = app.export_selected_xlsx([2, 5], job='scan-a', scope='approved', category=category)
        self.addCleanup(response.path.unlink, missing_ok=True)
        self.assertEqual(len(list(load_workbook(response.path, read_only=True).active.values)), 2)
        for response in (app.export_xlsx(job='scan-a', scope='approved', category='Not a real category'),
                         app.export_selected_xlsx([2], job='scan-a', scope='approved',
                                                  category='Not a real category')):
            self.assertEqual(response.status_code, 400)

    def test_category_exported_for_pending_and_approved_rows(self):
        with sqlite3.connect(app.DB_PATH) as conn:
            conn.execute("update reviews set review_category = 'Trademark Infringement Media' where id in (1, 2)")
        path = app.export_reviews_xlsx(job_id='scan-a')
        self.addCleanup(path.unlink, missing_ok=True)
        workbook = load_workbook(path, read_only=True)
        try:
            rows = list(workbook.active.values)
        finally:
            workbook.close()
        category_index = rows[0].index('Review Category')
        self.assertEqual(next(row[category_index] for row in rows[1:] if row[0] == 'approved'),
                         'Trademark Infringement Media')
        self.assertEqual(next(row[category_index] for row in rows[1:] if row[0] == 'pending'),
                         'Trademark Infringement Media')

    def test_saved_scan_negatives_survive_export(self):
        db = Path(self.folder.name) / "fresh.sqlite3"
        with patch.object(app, "DB_PATH", db):
            app.init_db()
            record = app.ListingRecord(
                job_id="new-scan", keyword="soap", search_url="", listing_url="https://example.com/1",
                seller_page_url="", title="soap", seller_name="seller", seller_id="seller-1",
                ship_from_country="Unknown", target_country="China", country_signal="Unknown",
                evidence_source="", evidence="", scan_status="match", negatives="tool, refill",
            )
            app.save_record(record)
            path = app.export_reviews_xlsx(job_id="new-scan", scope="match")
            self.addCleanup(path.unlink, missing_ok=True)
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(list(workbook.active.values)[1][:2], ("soap", "tool, refill"))
            finally:
                workbook.close()

    def test_false_positives_is_display_only_label(self):
        page = app.queue_page(state="not_queued", country_signal="all", job="scan-a").body.decode()
        self.assertIn('False Positives (1)', page)
        self.assertIn('state=not_queued', page)
        self.assertIn('value="not_queued" selected', page)
        self.assertNotIn('Not Queued', page)
        self.assertEqual(app.count_reviews('not_queued', job_id='scan-a'), 1)
        path = app.export_reviews_xlsx(job_id='scan-a', scope='not_queued')
        self.addCleanup(path.unlink, missing_ok=True)
        workbook = load_workbook(path, read_only=True)
        try:
            rows = list(workbook.active.values)
            self.assertEqual(rows[1][rows[0].index('Review Status')], 'False Positives')
        finally:
            workbook.close()
        script = (Path(__file__).parent / 'static' / 'fast_review.js').read_text()
        self.assertIn("not_queued: 'False Positives'", script)
        self.assertIn("pending: 'Pending'", script)

    def test_queue_page_has_export_choices(self):
        page = app.queue_page(state="approved", country_signal="all", job="scan-a").body.decode()
        self.assertIn('value="approved" selected', page)
        self.assertIn('value="match"', page)
        self.assertIn('id="export-selected"', page)
        self.assertIn('name="job" value="scan-a"', page)
        self.assertIn('id="export-category" name="category"', page)
        self.assertIn('<option value="">All categories</option>', page)
        self.assertIn('IP Rights Infringement Utility Outdoor &amp; Maintenence', page)

    def test_multi_scan_export_merges_ticked_scans_with_traceable_columns(self):
        with sqlite3.connect(app.DB_PATH) as conn:
            conn.executemany("insert into scan_history values (?, ?, null, 'done', '{}')",
                             [("scan-a", "2026-10-01T09:00:00+00:00"), ("scan-b", "2026-10-05T09:00:00+00:00")])
            conn.execute("insert into reviews (id,job_id,scan_status,review_state,country_signal,keyword,title,listing_url) "
                         "values (6, 'scan-c', 'match', 'approved', 'China', 'not ticked', 'item-6', 'url-6')")
        response = app.export_scans_xlsx(jobs=["scan-a", "scan-b"], scope="approved", category="")
        self.addCleanup(response.path.unlink, missing_ok=True)
        sheet = list(load_workbook(response.path, read_only=True).active.values)
        header, rows = sheet[0], sheet[1:]
        self.assertEqual(header[-2:], ("Scan ID", "Scan Date"))
        self.assertEqual(sorted((r[0], r[-2], r[-1]) for r in rows),
                         [("approved", "scan-a", "2026-10-01"), ("other scan", "scan-b", "2026-10-05")])
        self.assertEqual(response.filename, "prescreener_2_scans_approved.xlsx")
        self.assertNotIn("Scan ID", list(load_workbook(app.export_reviews_xlsx("scan-a"), read_only=True).active.values)[0])
        self.assertEqual(app.export_scans_xlsx(jobs=[], scope="all", category="").status_code, 400)
        self.assertEqual(app.export_scans_xlsx(jobs=["scan-a"], scope="all", category="Nope").status_code, 400)

    def test_history_page_offers_multi_scan_export(self):
        with sqlite3.connect(app.DB_PATH) as conn:
            conn.execute("alter table reviews add column market text")
            conn.execute("alter table reviews add column target_country text")
        page = app.history_page().body.decode()
        for text in ('action="/export-scans.xlsx"', 'type="date" id="range-from"',
                     'name="jobs" value="scan-a" form="multi-export"', 'id="export-category" name="category"',
                     'history_export.js'):
            with self.subTest(text=text):
                self.assertIn(text, page)


if __name__ == "__main__":
    unittest.main()
