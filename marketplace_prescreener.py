from __future__ import annotations

import json
import re
import ssl
import sqlite3
import sys
import tempfile
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from html import escape as html_escape
from pathlib import Path
from typing import Iterable, Literal
from urllib.error import HTTPError, URLError
from urllib.parse import quote_plus, urljoin, urlparse, parse_qs, urlencode, urlunparse
from urllib.request import Request, urlopen

from fastapi import Body, FastAPI, Query
from starlette.background import BackgroundTask
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.gzip import GZipMiddleware
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import decision_memory
import scan_history
import scan_history_page
from brand_signals import is_generic_brand
from counterfeit_signals import count_counterfeit_reviews_in_html, count_counterfeit_reviews_across_pages
from fuzzy_match import DEFAULT_STRICTNESS, SHORT_WORD_MAX, KeywordMatch, allowed_edits, disguised_spellings, keyword_words, legacy_contains, match_keyword
from ui_shell import STATIC_DIR, render_page

APP_TITLE = "Marketplace Country Prescreener"
USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
)
DEFAULT_COUNTRY = "China"
COUNTRY_OPTIONS = ["China", "US", "Any"]
COUNTRY_OPTION_LABELS = {"China": "China", "US": "US", "Any": "Any (no country filter)"}
DEFAULT_LIMIT = 0


@dataclass(frozen=True)
class Market:
    """A Walmart-banner storefront the scanner can point at.

    `domain` is the storefront host, `locale_prefix` is the path segment
    Walmart puts before `/ip/` and `/search` on that storefront (empty for
    walmart.com, "en" for walmart.ca, etc.), and `currency_suffix` is a
    cosmetic label appended to parsed prices so a $19.99 CAD listing isn't
    mistaken for USD.
    """

    code: str
    label: str
    domain: str
    locale_prefix: str = ""
    currency_suffix: str = ""

    @property
    def root_domain(self) -> str:
        return self.domain[4:] if self.domain.startswith("www.") else self.domain

    @property
    def path_prefix(self) -> str:
        return f"/{self.locale_prefix}" if self.locale_prefix else ""

    @property
    def base_url(self) -> str:
        return f"https://{self.domain}"


# NOTE on MX/CL domains: verified live via curl during setup (2026-08-25) for
# US and CA only. This sandbox's DNS resolver blocks .mx/.cl TLD lookups
# entirely (also blocked amazon.com.mx, a domain that definitely exists), so
# walmart.com.mx / lider.cl could not be confirmed from here. Both are the
# correct storefronts per public knowledge, but test them for real once this
# is running on a normal network before trusting MX/CL results.
MARKETS: dict[str, Market] = {
    "US": Market("US", "United States \u2014 walmart.com", "www.walmart.com"),
    "CA": Market("CA", "Canada \u2014 walmart.ca", "www.walmart.ca", locale_prefix="en", currency_suffix=" CAD"),
    "MX": Market("MX", "Mexico \u2014 walmart.com.mx (verify DNS resolves before use)", "www.walmart.com.mx", currency_suffix=" MXN"),
    "CL": Market("CL", "Chile \u2014 lider.cl (verify DNS resolves before use)", "www.lider.cl", currency_suffix=" CLP"),
}
DEFAULT_MARKET = "US"


def get_market(code: str | None) -> Market:
    return MARKETS.get((code or "").upper(), MARKETS[DEFAULT_MARKET])


SEARCH_TIMEOUT = 15
MAX_PAGE_SAFETY = 250
LISTING_WORKERS = 5
KEYWORD_WORKERS = 3
MAX_LISTING_WORKERS = 12
DEFAULT_STOP_AFTER_FIRST_MATCH = False
DEFAULT_MODE = "balanced"
DEFAULT_SEARCH_TYPE = "keyword"
DB_PATH = Path(__file__).with_name("prescreener.sqlite3")


class BotChallengeDetected(RuntimeError):
    pass


BOT_CHALLENGE_PATTERNS = [
    re.compile(r"not a robot", re.I),
    re.compile(r"are you human", re.I),
    re.compile(r"verify you are human", re.I),
    re.compile(r"robot check", re.I),
    re.compile(r"unusual traffic", re.I),
    re.compile(r"access denied", re.I),
    re.compile(r"pardon our interruption", re.I),
    re.compile(r"security check", re.I),
]


app = FastAPI(title=APP_TITLE)
app.add_middleware(GZipMiddleware, minimum_size=1024)  # queue JSON can be several MB
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
lock = threading.RLock()
JOBS: dict[str, dict] = {}
LATEST_JOB_ID: str | None = None

# Live jobs keep only a bounded recent window of scanned/results/logs in memory
# so a long-running scan (thousands of listings) doesn't grow an ever-larger
# JSON payload that gets re-shipped to the browser on every poll -- that was
# crashing extended-run scans. Every record is still saved to SQLite in full;
# this cap only affects the live "Latest scan rows" view, and accurate totals
# are tracked separately via job["stats"] regardless of trimming.
MAX_LIVE_SCANNED = 300
MAX_LIVE_RESULTS = 300
MAX_LIVE_LOGS = 300


@dataclass
class ListingRecord:
    job_id: str
    keyword: str
    search_url: str
    listing_url: str
    seller_page_url: str
    title: str
    seller_name: str
    seller_id: str
    ship_from_country: str
    target_country: str
    country_signal: str
    evidence_source: str
    evidence: str
    scan_status: str
    review_state: str = "pending"
    created_at: str = ""
    item_id: str = ""
    image_url: str = ""
    price_text: str = ""
    price_value: float | None = None
    description: str = ""
    counterfeit_review_count: int = 0
    market: str = DEFAULT_MARKET
    listing_brand: str = ""
    counterfeit_checked: bool = False
    match_kind: str = ""   # "" = exact, else "disguised" | "fuzzy"
    remembered: str = ""   # decision_memory note: "" | "auto_disputed" | "approved_before"
    negatives: str = ""    # comma-separated exclusions active when this row was scanned
    id: int | None = None


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def walmart_search_url(keyword: str, market: Market | None = None) -> str:
    market = market or MARKETS[DEFAULT_MARKET]
    return f"{market.base_url}{market.path_prefix}/search?q={quote_plus(keyword)}"


def is_bot_challenge(html: str) -> bool:
    text = normalize_space(strip_tags(html)).lower()
    title = page_title(html, fallback="").lower()
    combined = f"{title} {text}"
    return any(pattern.search(combined) for pattern in BOT_CHALLENGE_PATTERNS)


def fetch_html(url: str, timeout: int = SEARCH_TIMEOUT) -> str:
    request = Request(url, headers={"User-Agent": USER_AGENT})
    try:
        response = urlopen(request, timeout=timeout)
    except URLError as exc:
        # Some launch environments set SSL_CERT_FILE to a bundle missing the
        # public issuer. On macOS, retry with the OS CA bundle; never disable TLS.
        system_ca = Path("/etc/ssl/cert.pem")
        if (sys.platform != "darwin" or not system_ca.is_file()
                or not isinstance(exc.reason, ssl.SSLCertVerificationError)):
            raise
        context = ssl.create_default_context(cafile=str(system_ca))
        response = urlopen(request, timeout=timeout, context=context)
    with response:
        charset = response.headers.get_content_charset() or "utf-8"
        html = response.read().decode(charset, errors="replace")
    if is_bot_challenge(html):
        raise BotChallengeDetected(f"Bot challenge detected at {url}")
    return html


def strip_tags(text: str) -> str:
    return re.sub(r"<[^>]+>", " ", text)


def normalize_space(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def page_title(html: str, fallback: str = "") -> str:
    m = re.search(r"<title[^>]*>(.*?)</title>", html, re.I | re.S)
    return normalize_space(strip_tags(m.group(1))) if m else fallback


def extract_listing_urls(search_html: str, market: Market | None = None) -> list[str]:
    market = market or MARKETS[DEFAULT_MARKET]
    prefix = re.escape(market.path_prefix)
    domain = re.escape(market.domain)
    relative_pattern = rf"{prefix}/ip/[^\"'<>\s]+" if prefix else r"/ip/[^\"'<>\s]+"
    patterns = [rf"https://{domain}{prefix}/ip/[^\"'<>\s]+", relative_pattern]
    urls: list[str] = []
    seen: set[str] = set()
    for pattern in patterns:
        for raw in re.findall(pattern, search_html):
            url = raw if raw.startswith("http") else urljoin(market.base_url, raw)
            url = url.split("?")[0]
            if url not in seen:
                seen.add(url)
                urls.append(url)
    return urls


def extract_next_search_url(search_html: str, current_url: str) -> str:
    max_page_match = re.search(r'"paginationV2":\{"maxPage":(\d+)', search_html)
    max_page = int(max_page_match.group(1)) if max_page_match else None
    parsed = urlparse(current_url)
    params = parse_qs(parsed.query)
    current_page = 1
    if "page" in params and params["page"]:
        try:
            current_page = max(1, int(params["page"][0]))
        except ValueError:
            current_page = 1
    if max_page is not None and current_page >= max_page:
        return ""
    params["page"] = [str(current_page + 1)]
    next_query = urlencode(params, doseq=True)
    return urlunparse(parsed._replace(query=next_query))


def is_walmart_sold_and_shipped(product_html: str) -> bool:
    # Walmart embeds sellerName in the product page JSON. When the item is
    # sold and shipped directly by Walmart, this value is literally
    # "Walmart.com" -- verified against known Walmart-fulfilled and known
    # real third-party (China-based) listings, so this is a dedicated,
    # reliable signal rather than a fuzzy text guess.
    names = set(re.findall(r'"sellerName":"([^"]*)"', product_html))
    return names == {"Walmart.com"}


def extract_seller_page_url(product_html: str, listing_url: str) -> str:
    """Find the seller/store link on a product page.

    Domain-checked against `listing_url`'s own host (not a hardcoded
    "walmart.com") so this works for any configured Market -- walmart.ca,
    walmart.com.mx, lider.cl, etc. -- without needing the Market object
    threaded in here too.
    """
    patterns = [
        r'href=["\']([^"\']*(?:/seller/|seller\?|seller/|/store/|store/)[^"\']*)["\']',
        r'href=["\']([^"\']*seller[^"\']*)["\']',
    ]
    listing_domain = urlparse(listing_url).netloc
    domain_root = listing_domain[4:] if listing_domain.startswith("www.") else listing_domain
    candidates: list[str] = []
    seen: set[str] = set()
    for pattern in patterns:
        for raw in re.findall(pattern, product_html, flags=re.I):
            url = raw if raw.startswith("http") else urljoin(listing_url, raw)
            url = url.split("#")[0].split("?")[0]
            if url not in seen and domain_root and domain_root in url and "/ip/" not in url and "/search" not in url:
                seen.add(url)
                candidates.append(url)
    return candidates[0] if candidates else ""


def extract_seller_id(seller_page_url: str) -> str:
    if not seller_page_url:
        return "Unknown"
    from urllib.parse import urlparse, parse_qs
    parsed = urlparse(seller_page_url)
    params = parse_qs(parsed.query)
    for key in ("seller_id", "sellerId", "storeId", "merchantId", "id"):
        values = params.get(key)
        if values and values[0]:
            return normalize_space(values[0])
    segments = [seg for seg in parsed.path.split("/") if seg]
    for idx, seg in enumerate(segments):
        if seg.lower() in {"seller", "store", "merchant"} and idx + 1 < len(segments):
            candidate = segments[idx + 1]
            if re.fullmatch(r"\d{3,}", candidate):
                return candidate
    for seg in reversed(segments):
        if re.fullmatch(r"\d{3,}", seg):
            return seg
    m = re.search(r"(?:seller|store|merchant)[^\d]{0,20}(\d{3,})", seller_page_url, re.I)
    if m:
        return m.group(1)
    return "Unknown"


def extract_embedded_seller_id(product_html: str) -> str:
    """Fallback seller-id lookup for storefronts that don't render a clickable
    seller-page link (confirmed on walmart.ca product pages, which show a
    plain "Sold by X" text with no href, unlike walmart.com's linked seller
    storefronts). Walmart's product page still embeds the raw sellerId in a
    JSON payload regardless of storefront, so this pulls it from there.
    """
    m = re.search(r'"sellerId"\s*:\s*"([^"]+)"', product_html)
    return normalize_space(m.group(1)) if m else "Unknown"


def extract_item_id(listing_url: str) -> str:
    m = re.search(r"/ip/[^/]+/(\d+)", listing_url)
    if m:
        return m.group(1)
    tail = listing_url.rstrip("/").rsplit("/", 1)[-1]
    return tail if tail.isdigit() else "Unknown"


def extract_image_url(product_html: str) -> str:
    m = re.search(r'property=["\']og:image["\'][^>]*content=["\']([^"\']+)["\']', product_html, re.I)
    if m:
        return m.group(1)
    m = re.search(r'content=["\']([^"\']+)["\'][^>]*property=["\']og:image["\']', product_html, re.I)
    if m:
        return m.group(1)
    m = re.search(r'https://i\d*\.walmartimages\.[a-z.]+/[^"\'\\\s]+', product_html)
    if m:
        return m.group(0)
    return ""


def extract_description(product_html: str) -> str:
    m = re.search(r'property=["\']og:description["\'][^>]*content=["\']([^"\']*)["\']', product_html, re.I)
    if m:
        return normalize_space(strip_tags(m.group(1)))
    m = re.search(r'name=["\']description["\'][^>]*content=["\']([^"\']*)["\']', product_html, re.I)
    if m:
        return normalize_space(strip_tags(m.group(1)))
    m = re.search(r'"shortDescription"\s*:\s*"((?:[^"\\]|\\.)*)"', product_html)
    if m:
        text = m.group(1).encode().decode("unicode_escape", errors="ignore")
        return normalize_space(strip_tags(text))
    return ""


def extract_price_text(product_html: str) -> str:
    m = re.search(r'"currentPrice"\s*:\s*\{[^}]*?"price"\s*:\s*([0-9]+(?:\.[0-9]+)?)', product_html)
    if m:
        return f"${float(m.group(1)):.2f}"
    m = re.search(r'itemprop=["\']price["\'][^>]*content=["\']([0-9]+(?:\.[0-9]+)?)["\']', product_html, re.I)
    if m:
        return f"${float(m.group(1)):.2f}"
    m = re.search(r"\$\s?([0-9]{1,4}\.[0-9]{2})", product_html)
    if m:
        return f"${m.group(1)}"
    return "Unknown"


def extract_listing_brand(product_html: str) -> str:
    """Pull the seller-supplied "Listing Brand" attribute off a product page.

    Confirmed live (2026-08-27) that Walmart embeds this as a plain
    `"brand":"..."` JSON field on both search-result tiles and the product
    page itself, on both walmart.com and walmart.ca -- same field, same
    shape, no market-specific handling needed.
    """
    m = re.search(r'"brand"\s*:\s*"([^"]*)"', product_html)
    return normalize_space(m.group(1)) if m and m.group(1).strip() else ""


def parse_price_value(price_text: str) -> float | None:
    m = re.search(r"([0-9]+(?:\.[0-9]+)?)", price_text or "")
    if not m:
        return None
    try:
        return float(m.group(1))
    except ValueError:
        return None


def keyword_matches_text(keyword: str, text: str) -> bool:
    """Plain exact keyword-in-text check (no disguise decoding)."""
    return bool(text) and legacy_contains(keyword, text)


def find_keyword_match(keyword: str, title: str, description: str, match_description: bool, fuzzy_strictness: int | None = None) -> KeywordMatch | None:
    """Match a keyword against the title, optionally also the description.

    Disguised spellings (F3breze, F.e.b.r.e.z.e, Cyrillic lookalikes ...) are
    ALWAYS decoded; `fuzzy_strictness` (0-100, None = off) adds typo slack.
    With match_description on, EITHER the title OR the description may match.
    """
    found = match_keyword(keyword, title, fuzzy_strictness)
    if found is None and match_description and description:
        found = match_keyword(keyword, description, fuzzy_strictness)
        if found is not None:
            found = KeywordMatch(found.kind, found.details, "description")
    return found


def keyword_matches_listing(keyword: str, title: str, description: str, match_description: bool, fuzzy_strictness: int | None = None) -> bool:
    return find_keyword_match(keyword, title, description, match_description, fuzzy_strictness) is not None


def brand_matches(target_brand: str, listing_brand: str) -> bool:
    """Exact, case-insensitive match against the listing's own Brand attribute.

    Deliberately NOT a text-contains check against title/description -- live
    testing showed listing brand codes (e.g. a private bulk-seller brand like
    "WLBXH") almost never appear in the product title itself, so trusting the
    embedded brand field is the only reliable signal for this search type.
    """
    if not target_brand.strip():
        return True
    return listing_brand.strip().lower() == target_brand.strip().lower()


def term_matches_listing(
    search_type: str,
    term: str,
    title: str,
    description: str,
    listing_brand: str,
    match_description: bool,
    generic_only: bool = False,
    fuzzy_strictness: int | None = None,
) -> bool:
    """Dispatch to the right match rule for the job's search type.

    Brand-mode jobs use `brand_matches` (exact Brand attribute). Keyword mode
    always sees through disguises, adds typo slack when `fuzzy_strictness` is
    set, and with `generic_only` also requires a generic placeholder Brand.
    """
    if search_type == "brand":
        return brand_matches(term, listing_brand)
    if generic_only and not is_generic_brand(listing_brand):
        return False
    return keyword_matches_listing(term, title, description, match_description, fuzzy_strictness)


def country_matches_listing(target_country: str, country_signal: str, listing_brand: str,
                            search_type: str, generic_only: bool, generic_include: bool) -> bool:
    """Inclusive mode adds generic-brand hits without dropping ordinary country hits."""
    return (target_country == "Any" or generic_only or country_signal == target_country
            or (generic_include and search_type == "keyword" and is_generic_brand(listing_brand)))


def find_excluded_term(title: str, negative_terms: list[str] | tuple[str, ...]) -> str:
    """Return the first negative term found in the listing title, or "" if none.

    Words match at a word START, so 'tool' catches tool / tools / toolkit but
    not 'stool'. A multi-word term ('scraper blade') needs every word present,
    mirroring how positive keywords match. \\w keeps accented ca/mx/cl titles intact.
    """
    if not title:
        return ""
    text = title.lower()
    for term in negative_terms:
        words = re.findall(r"\w+", term.lower())
        if words and all(re.search(rf"(?<!\w){re.escape(w)}", text) for w in words):
            return term
    return ""


COUNTRY_HINTS: dict[str, list[re.Pattern[str]]] = {
    "China": [
        re.compile(r"(?:ships?|shipping|ship to|seller(?:\s+location)?|seller address|address|location|country of origin|based in)\s*[:\-]?\s*[^<>{}\n]{0,180}\b(?:china|cn)\b", re.I),
        re.compile(r"(?:ships?|shipping|ship to|seller(?:\s+location)?|seller address|address|location|country of origin|based in)[^<>{}\n]{0,180}\bchina\b", re.I),
        re.compile(r"(?:ships?|shipping|ship to|seller(?:\s+location)?|seller address|address|location|country of origin|based in)[^<>{}\n]{0,180}\bcn\b", re.I),
        re.compile(r"(?:ships?|shipping|ship to|seller(?:\s+location)?|seller address|address|location|country of origin|based in)[^<>{}\n]{0,220}\b(?:guangzhou|shenzhen|hangzhou|dongguan|foshan|yiwu|ningbo|suzhou|wenzhou|xiamen|qingdao|tianjin|beijing|shanghai|zhuhai|guangdong|zhejiang|jiangsu|shandong)\b", re.I),
        re.compile(r"(?:ships?|shipping|ship to|seller(?:\s+location)?|seller address|address|location|country of origin|based in)[^<>{}\n]{0,220}\b(?:GD|ZJ|JS|SH|BJ|CQ|SC|HN|HB|SD|FJ|JX|AH|GX|LN|TJ|HE|HL|JL)\s+\d{6}\b", re.I),
    ],
    "US": [
        re.compile(r"(?:ships?|shipping|ship to|seller(?:\s+location)?|seller address|address|location|based in)[^<>{}\n]{0,180}\b(?:united states|usa|us)\b", re.I),
        re.compile(r"(?:ships?|shipping|ship to|seller(?:\s+location)?|seller address|address|location|based in)[^<>{}\n]{0,220}\b[A-Z]{2}\s+\d{5}(?:-\d{4})?\b[^<>{}\n]{0,40}\bUS\b", re.I),
    ],
}


def first_match(text: str, patterns: Iterable[re.Pattern[str]]) -> str:
    for pattern in patterns:
        m = pattern.search(text)
        if m:
            if m.lastindex:
                return normalize_space(m.group(1))
            return normalize_space(m.group(0))
    return "Unknown"


def strip_review_text(html: str) -> str:
    """Remove embedded customer review text/title JSON fields before signal scanning.

    Walmart embeds full customer reviews directly in product/seller page HTML.
    Reviews are personal opinions/guesses ("I suspect it came from China"), not
    authoritative seller data, so they must not be able to trigger a country
    signal match the way real "Seller Location: ..." text would.
    """
    html = re.sub(r'"reviewText"\s*:\s*"(?:[^"\\]|\\.)*"', '"reviewText":""', html)
    html = re.sub(r'"reviewTitle"\s*:\s*"(?:[^"\\]|\\.)*"', '"reviewTitle":""', html)
    # Walmart also renders the review body as plain visible text (not just JSON),
    # consistently wrapped in this span structure. Strip that rendered copy too.
    html = re.sub(r'(<span class="tl-m db-m"><b></b>).*?(</span>)', r'\1\2', html, flags=re.S)
    return html


def extract_signals(html: str, target_country: str) -> tuple[str, str, str, str]:
    html = strip_review_text(html)
    text = normalize_space(strip_tags(html))
    china_evidence = first_match(text, COUNTRY_HINTS.get("China", []))
    us_evidence = first_match(text, COUNTRY_HINTS.get("US", []))

    if target_country == "China":
        if china_evidence != "Unknown":
            country_signal = "China"
        elif us_evidence != "Unknown":
            country_signal = "US"
        else:
            country_signal = "Unknown"
    elif target_country == "US":
        if us_evidence != "Unknown":
            country_signal = "US"
        elif china_evidence != "Unknown":
            country_signal = "China"
        else:
            country_signal = "Unknown"
    elif target_country == "Any":
        # No country filter requested -- still surface whichever signal (if
        # any) is present on the page for visibility, just don't gate on it.
        if china_evidence != "Unknown":
            country_signal = "China"
        elif us_evidence != "Unknown":
            country_signal = "US"
        else:
            country_signal = "Unknown"
    else:
        target_patterns = COUNTRY_HINTS.get(target_country, [re.compile(re.escape(target_country), re.I)])
        target_evidence = first_match(text, target_patterns)
        country_signal = target_country if target_evidence != "Unknown" else "Unknown"

    seller_name = first_match(
        text,
        [
            re.compile(r"seller name[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"sold by[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"visit the\s+([A-Za-z0-9&'\-., ]{2,80})\s+store", re.I),
            re.compile(r"seller:\s+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"by\s+([A-Za-z0-9&'\-., ]{2,80})", re.I),
        ],
    )

    ship_from = first_match(
        text,
        [
            re.compile(r"ships? from[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"shipping from[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"seller(?:\s+location)?[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"seller address[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"address[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"location[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"country of origin[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"based in[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
            re.compile(r"from[:\s]+([A-Za-z0-9&'\-., ]{2,80})", re.I),
        ],
    )

    if ship_from == "Unknown":
        if country_signal == "China" and china_evidence != "Unknown":
            ship_from = china_evidence
        elif country_signal == "US" and us_evidence != "Unknown":
            ship_from = us_evidence

    if country_signal == "China" and china_evidence != "Unknown":
        evidence = china_evidence
    elif country_signal == "US" and us_evidence != "Unknown":
        evidence = us_evidence
    elif ship_from != "Unknown":
        evidence = ship_from
    else:
        evidence = seller_name if seller_name != "Unknown" else "No explicit seller-country signal found on page"
    return country_signal, seller_name, ship_from, evidence


def init_db() -> None:
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """
            create table if not exists reviews (
                id integer primary key autoincrement,
                job_id text not null,
                keyword text not null,
                search_url text not null,
                listing_url text not null,
                seller_page_url text not null default '',
                title text not null,
                seller_name text not null,
                seller_id text not null default "Unknown",
                ship_from_country text not null,
                target_country text not null,
                country_signal text not null,
                evidence_source text not null default '',
                evidence text not null,
                scan_status text not null,
                review_state text not null,
                created_at text not null
            )
            """
        )
        cols = {row[1] for row in conn.execute("pragma table_info(reviews)").fetchall()}
        additions = {
            'seller_page_url': "alter table reviews add column seller_page_url text not null default ''",
            'seller_id': "alter table reviews add column seller_id text not null default 'Unknown'",
            'evidence_source': "alter table reviews add column evidence_source text not null default ''",
            'item_id': "alter table reviews add column item_id text not null default ''",
            'image_url': "alter table reviews add column image_url text not null default ''",
            'price_text': "alter table reviews add column price_text text not null default ''",
            'price_value': "alter table reviews add column price_value real",
            'description': "alter table reviews add column description text not null default ''",
            'counterfeit_review_count': "alter table reviews add column counterfeit_review_count integer not null default 0",
            'market': f"alter table reviews add column market text not null default '{DEFAULT_MARKET}'",
            'listing_brand': "alter table reviews add column listing_brand text not null default ''",
            'counterfeit_checked': "alter table reviews add column counterfeit_checked integer not null default 0",
            'match_kind': "alter table reviews add column match_kind text not null default ''",
            'negatives': "alter table reviews add column negatives text not null default ''",
        }
        for col, sql in additions.items():
            if col not in cols:
                conn.execute(sql)
        # Repair rows saved by the old shifted failure tuple (see _failed_scan):
        # status ended up in `evidence` and the message in `evidence_source`.
        # The match is exact and only hits blank-status rows, so this is a no-op
        # once repaired and safe on any copy of the DB (e.g. Ariel's).
        conn.execute(
            """
            update reviews
               set scan_status = evidence, evidence = evidence_source, evidence_source = 'product page'
             where scan_status = '' and evidence in ('error', 'blocked')
            """
        )
        decision_memory.init_memory(conn, extract_item_id)
        conn.commit()
    scan_history.init_history(DB_PATH)


init_db()


def save_record(record: ListingRecord) -> int:
    with sqlite3.connect(DB_PATH) as conn:
        # previously disputed items go straight to Disputed; approved ones get a note
        record.review_state, record.remembered, record.evidence = decision_memory.apply_to_new_row(
            conn, record.item_id, record.review_state, record.evidence)
        cur = conn.execute(
            """
            insert into reviews (
                job_id, keyword, search_url, listing_url, seller_page_url, title, seller_name, seller_id,
                ship_from_country, target_country, country_signal, evidence_source, evidence,
                scan_status, review_state, created_at, item_id, image_url, price_text, price_value, description, counterfeit_review_count, market, listing_brand, counterfeit_checked, match_kind, remembered, negatives
            ) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                record.job_id,
                record.keyword,
                record.search_url,
                record.listing_url,
                record.seller_page_url,
                record.title,
                record.seller_name,
                record.seller_id,
                record.ship_from_country,
                record.target_country,
                record.country_signal,
                record.evidence_source,
                record.evidence,
                record.scan_status,
                record.review_state,
                record.created_at,
                record.item_id,
                record.image_url,
                record.price_text,
                record.price_value,
                record.description,
                record.counterfeit_review_count,
                record.market,
                record.listing_brand,
                int(record.counterfeit_checked),
                record.match_kind,
                record.remembered,
                record.negatives,
            ),
        )
        conn.commit()
        return int(cur.lastrowid)


def _review_filters(review_state: str | None, country_signal: str | None, job_id: str | None) -> tuple[str, list[str]]:
    clauses: list[str] = []
    params: list[str] = []
    for column, value in (("review_state", review_state), ("country_signal", country_signal), ("job_id", job_id)):
        if value and value != "all":
            clauses.append(f"{column} = ?")
            params.append(value)
    return (" where " + " and ".join(clauses)) if clauses else "", params


def get_reviews(review_state: str | None = None, country_signal: str | None = None, job_id: str | None = None) -> list[dict]:
    where, params = _review_filters(review_state, country_signal, job_id)
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(f"select * from reviews{where} order by id desc", params).fetchall()
    return [dict(row) for row in rows]


def count_reviews(review_state: str | None = None, country_signal: str | None = None, job_id: str | None = None) -> int:
    where, params = _review_filters(review_state, country_signal, job_id)
    with sqlite3.connect(DB_PATH) as conn:
        return conn.execute(f"select count(*) from reviews{where}", params).fetchone()[0]


def update_review_state(record_id: int, review_state: str) -> None:
    update_review_states([record_id], review_state)


STATE_TAB_LABELS = {"pending": "Pending", "approved": "Approved", "rejected": "Disputed", "not_queued": "Not Queued"}
REMEMBERED_LABELS = {decision_memory.AUTO_DISPUTED: "Auto-disputed", decision_memory.APPROVED_BEFORE: "Approved before"}


REVIEW_STATES = ("pending", "approved", "rejected", "not_queued")


def update_review_states(record_ids: list[int], review_state: str) -> tuple[int, int]:
    """Set one review state on many rows in one transaction, remember the decision per
    item ID, and apply it to that item's other pending rows.
    Returns (rows changed, other rows changed by the memory cascade)."""
    if review_state not in REVIEW_STATES or not record_ids:
        return 0, 0
    with sqlite3.connect(DB_PATH) as conn:
        changed = 0
        for start in range(0, len(record_ids), 500):           # stay under SQLite's variable limit
            chunk = record_ids[start:start + 500]
            cur = conn.execute(f"update reviews set review_state = ? where id in ({','.join('?' * len(chunk))})",
                               (review_state, *chunk))
            changed += cur.rowcount
        cascaded = decision_memory.record_decisions(conn, record_ids, review_state, utc_now())
        conn.commit()
    return changed, cascaded


EXPORT_SCOPES = {
    "all": "Whole scan", "match": "Confirmed matches", "approved": "Approved",
    "pending": "Pending", "rejected": "Disputed", "not_queued": "Not queued",
}
ExportScope = Literal["all", "match", "approved", "pending", "rejected", "not_queued"]


def export_reviews_xlsx(job_id: str | None = None, scope: ExportScope = "all",
                        selected_ids: set[int] | None = None) -> Path:
    rows = get_reviews("all", job_id=job_id)
    if scope == "match":
        rows = [row for row in rows if row["scan_status"] == "match"]
    elif scope != "all":
        rows = [row for row in rows if row["review_state"] == scope]
    if selected_ids is not None:
        rows = [row for row in rows if row["id"] in selected_ids]
    any_counterfeit_checked = any(row.get("counterfeit_checked") for row in rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Queue"
    headers = ["Keyword", "Negatives", "Listing Brand", "Market", "Title", "Walmart URL", "Seller", "Seller ID", "Price", "Country", "Matched As", "Review Status", "Remembered"]
    if any_counterfeit_checked:
        headers.append("# of Counterfeit Reviews")
    ws.append(headers)
    for row in rows:
        values = [
            row.get("keyword", ""),
            row.get("negatives", ""),
            row.get("listing_brand", ""),
            row.get("market", DEFAULT_MARKET),
            row.get("title", ""),
            row.get("listing_url", ""),
            row.get("seller_name", "") or "Unknown",
            row.get("seller_id", "Unknown"),
            row.get("price_text", ""),
            row.get("country_signal", ""),
            (row.get("match_kind") or "exact").capitalize(),
            STATE_TAB_LABELS.get(row.get("review_state"), row.get("review_state", "")),
            REMEMBERED_LABELS.get(row.get("remembered") or "", ""),
        ]
        if any_counterfeit_checked:
            checked = row.get("counterfeit_checked")
            values.append(row.get("counterfeit_review_count", 0) if checked else "")
        ws.append(values)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"
    tmp = Path(tempfile.mkstemp(suffix=".xlsx", prefix="prescreener_")[1])
    wb.save(tmp)
    return tmp


def split_keywords(raw: str) -> list[str]:
    return [part.strip() for part in re.split(r"[\n,]+", raw or "") if part.strip()]


def create_job(keywords: list[str], country: str, limit: int, workers: int, mode: str, max_price: float | None = None, price_only: bool = False, match_description: bool = False, counterfeit_min: int | None = None, market: str = DEFAULT_MARKET, search_type: str = DEFAULT_SEARCH_TYPE, negative_terms: list[str] | None = None, generic_only: bool = False, fuzzy_strictness: int | None = None, disguised_search: bool = False, generic_include: bool = False) -> str:
    global LATEST_JOB_ID
    job_id = uuid.uuid4().hex[:10]
    market_obj = get_market(market)
    search_type = search_type if search_type in ("keyword", "brand") else DEFAULT_SEARCH_TYPE
    generic_only = bool(generic_only and search_type == "keyword")
    generic_include = bool(generic_include and search_type == "keyword" and not generic_only)
    fuzzy_strictness = fuzzy_strictness if search_type == "keyword" else None
    # (keyword it reports as, query typed into Walmart search). Disguised
    # spellings are extra searches whose listings still match the REAL keyword.
    search_plan = [[kw, kw] for kw in keywords]
    if disguised_search and search_type == "keyword":
        search_plan += [[kw, spelling] for kw in keywords for spelling in disguised_spellings(kw)]
    term_label = "listing brand" if search_type == "brand" else "keyword"
    job = {
        "id": job_id,
        "keywords": keywords,
        "country": country,
        "market": market_obj.code,
        "search_type": search_type,
        "limit": limit,
        "workers": workers,
        "mode": mode,
        "max_price": max_price,
        "price_only": bool(price_only and max_price is not None),
        "match_description": match_description,
        "negative_terms": list(negative_terms or []),
        "generic_only": generic_only,
        "generic_include": generic_include,
        "fuzzy_strictness": fuzzy_strictness,
        "search_plan": search_plan,
        "disguised_searches": len(search_plan) - len(keywords),
        "counterfeit_min": counterfeit_min,
        "status": "running",
        "message": "Queued",
        "current": 0,
        "total": max(len(search_plan), 1),
        "logs": [f"Queued scan for {len(keywords)} {term_label}(s) targeting {country} on {market_obj.label}."]
                + ([f"Excluding titles containing: {', '.join(negative_terms)}"] if negative_terms else [])
                + (["Generic listing brands only: matches also need a placeholder Brand (Generic, Unbranded, NOBRAND, Unknown, ...). Any country of origin counts; target country is ignored."] if generic_only else [])
                + (["Standard + generic: non-generic listings need the target country; generic listing brands match from any country."] if generic_include else [])
                + (["Disguise detection on: F3breze / F.e.b.r.e.z.e / lookalike letters count as the real keyword."] if search_type == "keyword" else [])
                + ([f"Fuzzy matching on at strictness {fuzzy_strictness}: " + "; ".join(f"{w} allows {allowed_edits(w, fuzzy_strictness)} letter difference(s)" for kw in keywords for w in keyword_words(kw))] if fuzzy_strictness is not None else [])
                + (["Also searching disguised spellings: " + ", ".join(q for kw, q in search_plan if q != kw)] if len(search_plan) > len(keywords) else []),
        "results": [],
        "scanned": [],
        "stats": {"scanned_total": 0, "results_total": 0, "china_total": 0, "us_total": 0, "unknown_total": 0, "excluded_total": 0},
        "stop_requested": False,
        "keyword_progress": {},
        "seen_urls": [],
        "created_at": utc_now(),
        "finished_at": None,
    }
    scan_history.save_scan(DB_PATH, job)
    with lock:
        JOBS[job_id] = job
        LATEST_JOB_ID = job_id
    return job_id


def append_log(job: dict, message: str) -> None:
    with lock:
        job["logs"].append(f"[{utc_now().split('T')[1][:8]}] {message}")
        if len(job["logs"]) > MAX_LIVE_LOGS:
            del job["logs"][: len(job["logs"]) - MAX_LIVE_LOGS]


def _failed_scan(item_id: str, status: str, evidence: str) -> tuple[str, str, str, str, str, str, str, str, str, str, str, str, str, str, int, str]:
    """scan_listing's result when the product page itself couldn't be loaded.

    One place to build it: the four failure paths used to hand-write this
    16-field tuple, dropped evidence_source, and every later field slid one
    slot left (the error text was saved as the source and the status blank).
    """
    return "Unknown", "", "", "", "Unknown", "Unknown", "Unknown", "product page", evidence, status, "", item_id, "", "Unknown", 0, ""


def scan_listing(keyword: str, listing_url: str, target_country: str, match_description: bool = False, check_counterfeit: bool = False, market: Market | None = None, search_type: str = DEFAULT_SEARCH_TYPE, negative_terms: tuple[str, ...] = (), generic_only: bool = False, fuzzy_strictness: int | None = None, generic_include: bool = False) -> tuple[str, str, str, str, str, str, str, str, str, str, str, str, str, str, int, str]:
    item_id = extract_item_id(listing_url)
    try:
        product_html = fetch_html(listing_url)
    except BotChallengeDetected as exc:
        return _failed_scan(item_id, "blocked", str(exc))
    except HTTPError as exc:
        return _failed_scan(item_id, "error", f"HTTP {exc.code}")
    except URLError as exc:
        return _failed_scan(item_id, "error", f"Network error: {exc.reason}")
    except Exception as exc:  # pragma: no cover
        return _failed_scan(item_id, "error", f"Unexpected error: {exc}")

    image_url = extract_image_url(product_html)
    price_text = extract_price_text(product_html)
    if price_text != "Unknown" and market and market.currency_suffix:
        price_text = f"{price_text}{market.currency_suffix}"
    title = page_title(product_html, fallback=listing_url.rsplit("/", 1)[-1])
    description = extract_description(product_html)
    listing_brand = extract_listing_brand(product_html)

    # Negative keywords: a title hit means the listing is out, full stop. Bail
    # before the seller-page and review fetches so excluded listings cost no
    # extra requests.
    excluded_by = find_excluded_term(title, negative_terms)
    if excluded_by:
        evidence = f"Title contains excluded term {excluded_by!r}"
        return "Unknown", title, description, "", "Unknown", "Unknown", "Unknown", "title", evidence, "excluded", "", item_id, image_url, price_text, 0, listing_brand

    # Sold-and-shipped-by-Walmart listings carry zero third-party/China
    # marketplace risk. Skip the seller-page lookup entirely for these --
    # saves a wasted request and keeps them out of match consideration.
    if is_walmart_sold_and_shipped(product_html):
        evidence = "Sold and shipped by Walmart.com (no third-party seller lookup needed)"
        return "US", title, description, "", "Walmart.com", "Walmart.com", "US", "product page", evidence, "walmart_sold", "", item_id, image_url, price_text, 0, listing_brand

    seller_page_url = extract_seller_page_url(product_html, listing_url)
    seller_page_title = ""
    seller_html = ""
    evidence_source = "product page"

    if seller_page_url:
        try:
            seller_html = fetch_html(seller_page_url)
            seller_page_title = page_title(seller_html, fallback=seller_page_url.rsplit("/", 1)[-1])
            evidence_source = "seller page"
        except BotChallengeDetected as exc:
            evidence_source = f"seller page blocked: {exc}"
        except Exception as exc:
            evidence_source = f"seller page fetch failed: {exc}"

    scan_html = seller_html or product_html
    country_signal, seller_name, ship_from, evidence = extract_signals(scan_html, target_country)
    seller_id = extract_seller_id(seller_page_url)
    if seller_id == "Unknown":
        seller_id = extract_embedded_seller_id(product_html)
    if check_counterfeit and seller_page_url and seller_html:
        counterfeit_review_count = count_counterfeit_reviews_across_pages(seller_page_url, seller_html, fetch_html)
    else:
        counterfeit_review_count = count_counterfeit_reviews_in_html(seller_html or product_html)


    if country_signal == "Unknown" and seller_html:
        fallback_country, fallback_seller, fallback_ship, fallback_evidence = extract_signals(product_html, target_country)
        if fallback_country != "Unknown":
            country_signal, seller_name, ship_from, evidence = fallback_country, fallback_seller, fallback_ship, fallback_evidence
            evidence_source = "product page fallback"

    if seller_name == "Unknown" and seller_page_title:
        seller_name = seller_page_title
    keyword_ok = term_matches_listing(search_type, keyword, title, description, listing_brand, match_description, generic_only, fuzzy_strictness)
    # "Any" means no country filter. Generic-only scans ignore country for
    # placeholders; inclusive scans add placeholders from any country while
    # still requiring the target country for every non-generic listing.
    country_ok = country_matches_listing(target_country, country_signal, listing_brand,
                                         search_type, generic_only, generic_include)
    if country_ok:
        evidence = f"{evidence_source}: {evidence}"
        if generic_only and keyword_ok:
            evidence = f"{evidence} (generic listing brand {listing_brand!r}; country of origin ignored for generic-only searches)"
        elif generic_include and keyword_ok and is_generic_brand(listing_brand) and country_signal != target_country:
            evidence = f"{evidence} (generic listing brand {listing_brand!r}; country of origin ignored for generic-brand match)"
        if not keyword_ok:
            if search_type == "brand":
                evidence = f"{evidence} (listing brand {listing_brand or 'Unknown'!r} does not match target brand {keyword!r}, treating as non-match)"
            elif generic_only and not is_generic_brand(listing_brand):
                evidence = f"{evidence} (listing brand {listing_brand or 'none'!r} is not a generic brand, treating as non-match)"
            else:
                where = "title or description" if match_description else "title"
                slack = f" or with up to the allowed letter differences (strictness {fuzzy_strictness})" if fuzzy_strictness is not None else ""
                evidence = f"{evidence} (keyword {keyword!r} not found in {where}, not even disguised{slack}, treating as non-match)"
    status = "match" if (country_ok and keyword_ok) else "unknown"
    return country_signal, title, description, seller_page_url, seller_name, seller_id, ship_from, evidence_source, evidence, status, seller_page_title, item_id, image_url, price_text, counterfeit_review_count, listing_brand


def scan_listing_job(keyword: str, listing_url: str, target_country: str, search_url: str, match_description: bool = False, check_counterfeit: bool = False, market: Market | None = None, search_type: str = DEFAULT_SEARCH_TYPE, negative_terms: tuple[str, ...] = (), generic_only: bool = False, fuzzy_strictness: int | None = None, generic_include: bool = False) -> tuple[ListingRecord, str]:
    market = market or MARKETS[DEFAULT_MARKET]
    country_signal, title, description, seller_page_url, seller_name, seller_id, ship_from, evidence_source, evidence, status, seller_page_title, item_id, image_url, price_text, counterfeit_review_count, listing_brand = scan_listing(keyword, listing_url, target_country, match_description, check_counterfeit, market, search_type, negative_terms, generic_only, fuzzy_strictness, generic_include)
    found = find_keyword_match(keyword, title, description, match_description, fuzzy_strictness) if search_type == "keyword" else None
    match_kind = found.kind if found and found.kind != "exact" else ""
    if match_kind:
        evidence = f"{evidence} [{match_kind} match: {found.describe()}]"
    record = ListingRecord(
        job_id="",
        keyword=keyword,
        search_url=search_url,
        listing_url=listing_url,
        seller_page_url=seller_page_url,
        title=title or seller_page_title or listing_url,
        description=description,
        seller_name=seller_name,
        seller_id=seller_id,
        ship_from_country=ship_from,
        target_country=target_country,
        country_signal=country_signal,
        evidence_source=evidence_source,
        evidence=evidence,
        scan_status=status,
        review_state="pending" if status == "match" else "not_queued",
        created_at=utc_now(),
        item_id=item_id,
        image_url=image_url,
        price_text=price_text,
        price_value=parse_price_value(price_text),
        counterfeit_review_count=counterfeit_review_count,
        market=market.code,
        listing_brand=listing_brand,
        counterfeit_checked=check_counterfeit,
        match_kind=match_kind,
    )
    return record, status


def run_job(job_id: str) -> None:
    with lock:
        job = JOBS[job_id]
    keywords = job["keywords"]
    country = job["country"]
    market = get_market(job.get("market"))
    search_type = job.get("search_type", DEFAULT_SEARCH_TYPE)
    page_limit = job["limit"]
    mode = job.get("mode", DEFAULT_MODE)
    max_price = job.get("max_price")
    price_only = bool(job.get("price_only") and max_price is not None)
    match_description = bool(job.get("match_description"))
    negative_terms = tuple(job.get("negative_terms") or ())
    generic_only = bool(job.get("generic_only"))
    generic_include = bool(job.get("generic_include"))
    fuzzy_strictness = job.get("fuzzy_strictness")
    search_plan = job.get("search_plan") or [[kw, kw] for kw in keywords]
    counterfeit_min = job.get("counterfeit_min")
    worker_cap = max(1, min(int(job.get("workers", LISTING_WORKERS)), MAX_LISTING_WORKERS))
    if mode == "fast":
        worker_cap = max(worker_cap, 8)
    elif mode == "thorough":
        worker_cap = min(worker_cap, 3)

    seen_urls: set[str] = set(job.get("seen_urls") or [])
    discovered = len(job.get("scanned") or [])

    def process_keyword(position: int, keyword: str, query: str) -> None:
        nonlocal discovered
        with lock:
            if job.get("stop_requested"):
                return
            progress = (job.get("keyword_progress") or {}).get(query)
            job["message"] = f"Searching {market.label} for {query!r} ({position}/{len(search_plan)})"
        if progress and progress.get("finished"):
            append_log(job, f"Skipping {query!r}: already finished in a previous run.")
            return
        append_log(job, job["message"])

        if progress and progress.get("next_url"):
            current_search_url = progress["next_url"]
            pages_crawled = progress.get("pages_crawled", 0)
            append_log(job, f"Resuming {query!r} from page {pages_crawled + 1}: {current_search_url}")
        else:
            current_search_url = walmart_search_url(query, market)
            pages_crawled = 0
        visited_search_pages: set[str] = set()
        keyword_listing_count = 0

        while current_search_url and current_search_url not in visited_search_pages:
            with lock:
                if job.get("stop_requested"):
                    job["status"] = "stopped"
                    job["message"] = "Scan stopped by user."
                    job["finished_at"] = utc_now()
                    job.setdefault("keyword_progress", {})[query] = {
                        "next_url": current_search_url,
                        "pages_crawled": pages_crawled,
                        "finished": False,
                    }
                    job["seen_urls"] = list(seen_urls)
                    append_log(job, "User requested stop. Ending scan early.")
                    return
            if page_limit > 0 and pages_crawled >= page_limit:
                append_log(job, f"Page cap reached for {query!r}: {page_limit} page(s).")
                break
            if pages_crawled >= MAX_PAGE_SAFETY:
                append_log(job, f"Safety cap reached for {query!r}: {MAX_PAGE_SAFETY} pages.")
                break

            visited_search_pages.add(current_search_url)
            pages_crawled += 1
            append_log(job, f"Fetching search page {pages_crawled} for {query!r}: {current_search_url}")

            try:
                search_html = fetch_html(current_search_url)
            except BotChallengeDetected as exc:
                append_log(job, f"Search page blocked for {query!r} on page {pages_crawled}: {exc}")
                break
            except Exception as exc:
                append_log(job, f"Search failed for {query!r} on page {pages_crawled}: {exc}")
                break

            page_listing_urls = extract_listing_urls(search_html, market)
            next_search_url = extract_next_search_url(search_html, current_search_url)
            append_log(job, f"Page {pages_crawled} yielded {len(page_listing_urls)} candidate listing(s).")

            with lock:
                new_urls: list[str] = []
                for u in page_listing_urls:
                    if u not in seen_urls:
                        seen_urls.add(u)
                        new_urls.append(u)
            if not new_urls:
                append_log(job, f"No new listings found on page {pages_crawled} for {query!r}.")
            else:
                worker_count = min(worker_cap, len(new_urls))
                append_log(job, f"Scanning {len(new_urls)} new listing(s) with {worker_count} worker(s).")
                # NOTE: managed manually (not via `with`) so a mid-batch stop can
                # cancel queued-but-not-yet-started listings immediately instead of
                # blocking until the *entire* page's listings finish scanning --
                # that block-until-done behavior is what used to leave jobs stuck
                # in "stopping" forever (Resume button never appeared).
                pool = ThreadPoolExecutor(max_workers=worker_count)
                futures = {
                    pool.submit(scan_listing_job, keyword, listing_url, country, current_search_url, match_description, counterfeit_min is not None, market, search_type, negative_terms, generic_only, fuzzy_strictness, generic_include): listing_url
                    for listing_url in new_urls
                }
                completed_urls: set[str] = set()
                stopped_mid_page = False
                for future in as_completed(futures):
                    listing_url = futures[future]
                    completed_urls.add(listing_url)
                    try:
                        record, status = future.result()
                    except Exception as exc:
                        append_log(job, f"Listing scan failed for {listing_url}: {exc}")
                    else:
                        with lock:
                            discovered += 1
                            keyword_listing_count += 1
                            job["message"] = f"Finished scan {discovered} for {query!r} on page {pages_crawled}"
                        record.job_id = job_id
                        record.negatives = ", ".join(negative_terms)
                        keyword_ok = term_matches_listing(search_type, keyword, record.title, record.description, record.listing_brand, match_description, generic_only, fuzzy_strictness)
                        if status in ("walmart_sold", "excluded"):
                            # Sold-and-shipped-by-Walmart carries zero third-party
                            # risk, and excluded titles were filtered out on purpose:
                            # never let price/counterfeit promotion flip either to a match.
                            pass
                        elif status in ("match", "unknown") and max_price is not None:
                            price_ok = record.price_value is not None and record.price_value <= max_price
                            if price_only:
                                status = "match" if (price_ok and keyword_ok) else "unknown"
                            elif status == "match" and not price_ok:
                                status = "unknown"
                            record.scan_status = status
                            record.review_state = "pending" if status == "match" else "not_queued"
                        if status not in ("match", "walmart_sold", "excluded") and keyword_ok and counterfeit_min is not None and record.counterfeit_review_count >= counterfeit_min:
                            status = "match"
                            record.scan_status = status
                            record.review_state = "pending"
                        record.id = save_record(record)
                        with lock:
                            stats = job["stats"]
                            stats["scanned_total"] += 1
                            if status == "excluded":
                                # filtered out on purpose; don't let them inflate the Unknown count
                                stats["excluded_total"] += 1
                            else:
                                signal_key = {"China": "china_total", "US": "us_total"}.get(record.country_signal, "unknown_total")
                                stats[signal_key] += 1
                            job["scanned"].append(asdict(record))
                            if len(job["scanned"]) > MAX_LIVE_SCANNED:
                                del job["scanned"][: len(job["scanned"]) - MAX_LIVE_SCANNED]
                            if status == "match":
                                stats["results_total"] += 1
                                job["results"].append(asdict(record))
                                if len(job["results"]) > MAX_LIVE_RESULTS:
                                    del job["results"][: len(job["results"]) - MAX_LIVE_RESULTS]
                        if status == "match":
                            if record.remembered == decision_memory.AUTO_DISPUTED:
                                append_log(job, f"Match auto-disputed (you disputed item {record.item_id} before): {record.title or listing_url}")
                            else:
                                append_log(job, f"Match found: {record.title or listing_url}")
                    with lock:
                        if job.get("stop_requested"):
                            stopped_mid_page = True
                    if stopped_mid_page:
                        break

                if stopped_mid_page:
                    # Cancel whatever hasn't started yet -- already-running threads
                    # finish on their own but we don't wait around for them.
                    pool.shutdown(wait=False, cancel_futures=True)
                    unprocessed = [u for u in new_urls if u not in completed_urls]
                    with lock:
                        for u in unprocessed:
                            seen_urls.discard(u)
                        job["status"] = "stopped"
                        job["message"] = "Scan stopped by user."
                        job["finished_at"] = utc_now()
                        job.setdefault("keyword_progress", {})[query] = {
                            "next_url": current_search_url,
                            "pages_crawled": pages_crawled - 1,
                            "finished": False,
                        }
                        job["seen_urls"] = list(seen_urls)
                    append_log(job, f"User requested stop mid-page ({len(unprocessed)} listing(s) on this page left for next resume). Ending scan early.")
                    return
                pool.shutdown(wait=True)

            if not next_search_url or next_search_url in visited_search_pages:
                append_log(job, f"No more search pages for {query!r}.")
                break
            current_search_url = next_search_url

        append_log(job, f"Finished keyword {query!r}: crawled {pages_crawled} page(s), inspected {keyword_listing_count} new listing(s), queued {job['stats']['results_total']} match(es) so far.")
        with lock:
            job["current"] = min(job["total"], job.get("current", 0) + 1)
            job.setdefault("keyword_progress", {})[query] = {"next_url": None, "pages_crawled": pages_crawled, "finished": True}
            job["seen_urls"] = list(seen_urls)

    try:
        keyword_worker_count = min(KEYWORD_WORKERS, len(search_plan)) if search_plan else 1
        append_log(job, f"Running up to {keyword_worker_count} keyword search(es) in parallel.")
        with ThreadPoolExecutor(max_workers=keyword_worker_count) as pool:
            futures = {
                pool.submit(process_keyword, i, keyword, query): query
                for i, (keyword, query) in enumerate(search_plan, start=1)
            }
            for future in as_completed(futures):
                keyword = futures[future]
                try:
                    future.result()
                except Exception as exc:
                    append_log(job, f"Keyword task failed for {keyword!r}: {exc}")

        with lock:
            if job.get("stop_requested"):
                job["status"] = "stopped"
                job["message"] = f"Stopped. {job['stats']['results_total']} match(es) queued for review."
            else:
                job["status"] = "done"
                job["message"] = f"Finished. {job['stats']['results_total']} match(es) queued for review."
            job["current"] = job["total"]
            job["finished_at"] = utc_now()
        append_log(job, job["message"])
    except Exception as exc:  # pragma: no cover
        with lock:
            job["status"] = "error"
            job["message"] = f"Job failed: {exc}"
            job["finished_at"] = utc_now()
        append_log(job, job["message"])
    finally:
        scan_history.update_status(DB_PATH, job)


RESIZABLE_TABLE_SCRIPT = """
function initResizableTables() {
  var tables = document.querySelectorAll('table[data-resizable]');
  for (var t = 0; t < tables.length; t++) {
    var table = tables[t];
    if (table.dataset.resizeReady === 'true') continue;
    table.dataset.resizeReady = 'true';
    table.style.tableLayout = 'fixed';
    var headers = table.querySelectorAll('thead th');
    for (var i = 0; i < headers.length; i++) {
      (function (th) {
        var handle = document.createElement('span');
        handle.className = 'col-resize-handle';
        if (getComputedStyle(th).position === 'static') {
          th.style.position = 'relative';
        }
        th.appendChild(handle);
        var startX = 0;
        var startWidth = 0;
        function onMove(e) {
          var delta = e.clientX - startX;
          var newWidth = Math.max(50, startWidth + delta);
          th.style.width = newWidth + 'px';
        }
        function onUp() {
          document.removeEventListener('pointermove', onMove);
          document.removeEventListener('pointerup', onUp);
          handle.classList.remove('active');
        }
        handle.addEventListener('pointerdown', function (e) {
          e.preventDefault();
          startX = e.clientX;
          startWidth = th.getBoundingClientRect().width;
          handle.classList.add('active');
          document.addEventListener('pointermove', onMove);
          document.addEventListener('pointerup', onUp);
        });
      })(headers[i]);
    }
  }
}
"""

THUMB_PREVIEW_SCRIPT = """
function initThumbPreview() {
  var preview = document.getElementById('thumb-preview');
  if (!preview) {
    preview = document.createElement('div');
    preview.id = 'thumb-preview';
    preview.className = 'thumb-preview';
    var img = document.createElement('img');
    preview.appendChild(img);
    document.body.appendChild(preview);
  }
  var previewImg = preview.querySelector('img');
  function findThumbLink(target) {
    while (target && target !== document.body) {
      if (target.classList && target.classList.contains('thumb-link')) return target;
      target = target.parentNode;
    }
    return null;
  }
  document.addEventListener('mouseover', function (e) {
    var link = findThumbLink(e.target);
    if (!link) return;
    var src = link.getAttribute('data-preview-src');
    if (!src) return;
    previewImg.src = src;
    preview.style.display = 'block';
  });
  document.addEventListener('mousemove', function (e) {
    if (preview.style.display !== 'block') return;
    var x = e.clientX + 15;
    var y = e.clientY + 15;
    var maxX = window.innerWidth - 320;
    var maxY = window.innerHeight - 320;
    if (x > maxX) x = maxX;
    if (y > maxY) y = maxY;
    preview.style.left = x + 'px';
    preview.style.top = y + 'px';
  });
  document.addEventListener('mouseout', function (e) {
    var link = findThumbLink(e.target);
    if (!link) return;
    var toLink = findThumbLink(e.relatedTarget);
    if (toLink) return;
    preview.style.display = 'none';
  });
}
"""

SEARCH_FORM_SCRIPT = r"""
function updateSearchTypeLabels() {
  var isBrand = document.querySelector('input[name="search_type"]:checked').value === 'brand';
  document.getElementById('keywords-label').textContent = isBrand ? 'Listing Brand(s)' : 'Keywords';
  document.getElementById('keywords').placeholder = isBrand
    ? 'Example:\nWLBXH\nSOMEOTHERBRAND'
    : 'Example:\ndish soap\nlaundry detergent\npaper towels';
  document.getElementById('keywords-hint').textContent = isBrand
    ? 'One listing brand per line (exact match, case-insensitive). Every listing tagged with that brand is pulled, regardless of what the title says.'
    : 'One per line (or comma-separated). Each keyword is searched, then listings are matched by keyword-in-title.';
  document.getElementById('generic-mode-wrap').hidden = isBrand;
  document.getElementById('match-options').hidden = isBrand;
}

function savePrescreenerPref(key, value) {
  localStorage.setItem('prescreener_' + key, value);
}

function restorePrescreenerPrefs() {
  var market = localStorage.getItem('prescreener_market');
  var marketSelect = document.getElementById('market');
  if (market && marketSelect) marketSelect.value = market;

  var searchType = localStorage.getItem('prescreener_search_type');
  if (searchType) {
    var radio = document.querySelector('input[name="search_type"][value="' + searchType + '"]');
    if (radio) radio.checked = true;
  }

  var country = localStorage.getItem('prescreener_country');
  var countrySelect = document.getElementById('country');
  if (country && countrySelect) countrySelect.value = country;

  var genericMode = document.getElementById('generic_mode');
  if (genericMode) {
    var savedMode = localStorage.getItem('prescreener_generic_mode');
    // Preserve the old checked toggle preference when moving to three choices.
    genericMode.value = ['standard', 'only', 'include'].includes(savedMode) ? savedMode :
      (localStorage.getItem('prescreener_generic_only') === 'true' ? 'only' : 'standard');
  }

  updateSearchTypeLabels();
}

document.addEventListener('DOMContentLoaded', restorePrescreenerPrefs);
"""


def html_page(title: str, body: str, extra_script: str = "", *, active: str = "", subtitle: str = "", scripts: tuple[str, ...] = ()) -> HTMLResponse:
    """All pages render through the ProductWatch-style shell (ui_shell.py + static/ui.css)."""
    return render_page(title, body, active=active, subtitle=subtitle, extra_script=extra_script,
                       latest_job_id=LATEST_JOB_ID or scan_history.latest_scan_id(DB_PATH),
                       doc_title=APP_TITLE, scripts=scripts)


@app.get("/", response_class=HTMLResponse)
def home() -> HTMLResponse:
    body = f"""
  <div class="card">
    <form action="/scan" method="get">
      <div class="grid" style="grid-template-columns: 1fr;">
        <div>
          <label>Search by</label>
          <div style="display:flex; gap:20px; margin-bottom:8px;">
            <label style="display:flex; align-items:center; gap:6px; margin-bottom:0; font-weight:400;">
              <input type="radio" name="search_type" value="keyword" checked onchange="savePrescreenerPref('search_type', this.value); updateSearchTypeLabels();" style="width:auto;"> Keyword
            </label>
            <label style="display:flex; align-items:center; gap:6px; margin-bottom:0; font-weight:400;">
              <input type="radio" name="search_type" value="brand" onchange="savePrescreenerPref('search_type', this.value); updateSearchTypeLabels();" style="width:auto;"> Listing Brand
            </label>
          </div>
          <label for="keywords" id="keywords-label">Keywords</label>
          <textarea id="keywords" name="keywords" placeholder="Example:\ndish soap\nlaundry detergent\npaper towels"></textarea>
          <p class="tiny" id="keywords-hint">One per line (or comma-separated). Each keyword is searched, then listings are matched by keyword-in-title.</p>
          <div id="generic-mode-wrap" style="margin-top:8px;">
            <label for="generic_mode">Listing brand matching</label>
            <select id="generic_mode" name="generic_mode" aria-describedby="generic-mode-hint" onchange="savePrescreenerPref('generic_mode', this.value)">
              <option value="standard">Standard: target-country listings</option>
              <option value="only">Generic brands only: any country</option>
              <option value="include">Standard + generic brands: any country for generic</option>
            </select>
            <p class="tiny" id="generic-mode-hint">All options require the keyword to match. Standard uses your target country; generic-only keeps just placeholder listing brands regardless of country. Standard + generic keeps regular target-country matches <strong>and</strong> adds placeholder listing brands from any country. Sold-and-shipped-by-Walmart listings remain excluded.</p>
          </div>
          <fieldset id="match-options" class="match-panel">
            <legend>How keywords are matched</legend>
            <p class="match-always" id="disguise-hint">
              <span class="badge kind-disguised">Always on</span>
              <strong>Disguise detection.</strong> Catches sellers hiding the keyword: numbers or symbols for letters (F3breze, Fe8reze, Febr@ze), lookalike letters from other alphabets, accents (F&eacute;br&egrave;ze), dots, dashes or spaces inside the word (F.e.b.r.e.z.e, Feb reze), invisible characters and stretched letters (Febreeeze). Disguises are never counted as typos, so this works whatever the settings below.
            </p>
            <label class="check-row">
              <input id="fuzzy" name="fuzzy" type="checkbox" value="true" aria-describedby="fuzzy-hint">
              Enable fuzzy matching (also catch typos and misspellings)
            </label>
            <p class="tiny" id="fuzzy-hint">Lets a few letters be wrong, missing, extra or swapped (Febrese, Febrez, Febreeze). Words of 4 letters or fewer never get typos, because they'd match too many other words (tide &rarr; tile, ride).</p>
            <div id="fuzzy-controls" class="fuzzy-controls" hidden>
              <label for="fuzzy_strictness">Strictness: <output id="fuzzy-strictness-value" for="fuzzy_strictness">{DEFAULT_STRICTNESS}</output></label>
              <input id="fuzzy_strictness" name="fuzzy_strictness" type="range" min="0" max="100" step="5" value="{DEFAULT_STRICTNESS}" aria-describedby="fuzzy-scale">
              <div id="fuzzy-scale" class="fuzzy-scale tiny">
                <span>0 &middot; anything goes</span><span>70 &middot; ~2 letters</span><span>90 &middot; ~1 letter</span><span>100 &middot; no typos</span>
              </div>
            </div>
            <div class="match-preview" id="match-preview" aria-live="polite"></div>
            <label class="check-row">
              <input id="disguised_search" name="disguised_search" type="checkbox" value="true" aria-describedby="disguised-search-hint">
              Also search disguised spellings (slower)
            </label>
            <p class="tiny" id="disguised-search-hint">Walmart's search usually won't show F3breze listings when you search &ldquo;febreze&rdquo;, so this also searches up to 5 disguised spellings of each keyword. Each spelling is a full extra search, so it's slower and hits more robot checks.</p>
          </fieldset>
          <label for="negative_keywords" style="margin-top:12px;">Exclude titles containing <span class="tiny">(optional)</span></label>
          <textarea id="negative_keywords" name="negative_keywords" style="min-height:72px;" aria-describedby="negative-hint" placeholder="Example:\ntool\nrefill"></textarea>
          <p class="tiny" id="negative-hint">One per line (or comma-separated). Any listing whose title contains one of these is skipped, even if it matches a keyword. Matches word starts, so &ldquo;tool&rdquo; also drops &ldquo;tools&rdquo; and &ldquo;toolkit&rdquo; but not &ldquo;stool&rdquo;. Applies to Listing Brand searches too.</p>
        </div>
      </div>
        <div class="grid" style="margin-top:12px;">
          <div>
            <label for="market">Market (which Walmart storefront to scan) &mdash; remembered between visits</label>
            <select id="market" name="market" onchange="savePrescreenerPref('market', this.value)">
              {''.join(f'<option value="{html_escape(m.code)}"{" selected" if m.code == DEFAULT_MARKET else ""}>{html_escape(m.label)}</option>' for m in MARKETS.values())}
            </select>
          </div>
        </div>
      <details class="settings-accordion" style="margin-top:12px;">
        <summary class="button secondary settings-toggle">Scan settings</summary>
        <div class="grid" style="margin-top:12px;">
          <div>
            <label for="country">Target country</label>
            <select id="country" name="country" onchange="savePrescreenerPref('country', this.value)">
              {''.join(f'<option value="{html_escape(option)}"{" selected" if option == DEFAULT_COUNTRY else ""}>{html_escape(COUNTRY_OPTION_LABELS.get(option, option))}</option>' for option in COUNTRY_OPTIONS)}
            </select>
          </div>
          <div>
            <label for="limit">Max search pages per keyword (0 = all)</label>
            <input id="limit" name="limit" type="number" min="0" max="500" value="{DEFAULT_LIMIT}">
          </div>
        </div>
        <div class="grid" style="margin-top:12px; grid-template-columns: .8fr .8fr 1.4fr;">
          <div>
            <label for="mode">Scan mode</label>
            <select id="mode" name="mode">
              <option value="fast">Fast</option>
              <option value="balanced" selected>Balanced</option>
              <option value="thorough">Thorough</option>
            </select>
          </div>
          <div>
            <label for="workers">Listing workers</label>
            <input id="workers" name="workers" type="number" min="1" max="12" value="{LISTING_WORKERS}">
          </div>
          <div style="align-self:end; padding-top:18px;"></div>
        </div>
        <div class="grid" style="margin-top:12px; grid-template-columns: .8fr 1.2fr;">
          <div>
            <label for="max_price">Max price ($, optional)</label>
            <input id="max_price" name="max_price" type="number" min="0" step="0.01" placeholder="No limit">
          </div>
          <div style="align-self:end;">
            <label style="display:flex; align-items:center; gap:8px; margin-bottom:0;">
              <input id="price_only" name="price_only" type="checkbox" value="true" style="width:auto;">
              Match by price only (ignore country)
            </label>
          </div>
        </div>
        <div class="grid" style="margin-top:12px; grid-template-columns: 1fr;">
          <div>
            <label style="display:flex; align-items:center; gap:8px; margin-bottom:0;">
              <input id="match_description" name="match_description" type="checkbox" value="true" style="width:auto;">
              Also match keyword against listing description (not just title)
            </label>
          </div>
        </div>
        <div class="grid" style="margin-top:12px; grid-template-columns: 1fr 1.5fr;">
          <div>
            <label for="counterfeit_min">Min. counterfeit-mentioning reviews (optional)</label>
            <input id="counterfeit_min" name="counterfeit_min" type="number" min="0" step="1" placeholder="Off">
          </div>
          <div style="align-self:end;">
            <span class="tiny">If set, a seller with at least this many reviews mentioning fake/counterfeit/knockoff language counts as a match too &mdash; in China, US, or Unknown, regardless of country signal.</span>
          </div>
        </div>
      </details>
      <div class="row-actions" style="margin-top:12px;">
        <button type="submit">Scan listings</button>
        <a class="button secondary" href="/queue">Review queue</a>
        <a class="button secondary" href="/queue">Choose export</a>
      </div>
    </form>
    <p class="muted" style="margin-top:12px;">
      What it does: searches the selected <strong>Market</strong> storefront for each keyword, crawls through search result pages until there are no more pages (or your page cap is reached), opens each listing, and flags pages that explicitly mention the target country.
      Pick <strong>Market</strong> to point the scan at walmart.com, walmart.ca, walmart.com.mx, or lider.cl &mdash; everything else (target country, price, keyword matching) behaves the same regardless of market.
      If Walmart does not show a seller-country clue, the listing stays unflagged or unknown. If a robot check or captcha pops up, the scan logs it and moves on.
      Open <strong>Scan settings</strong> to set a max price. With "Match by price only" checked, listings under the price count as matches regardless of country; unchecked, price is an extra filter on top of the country match.
      By default the keyword must appear in the listing title; check "Also match against description" to count a listing as a match if the keyword shows up in EITHER the title or the description.
      Keyword searches <strong>always see through disguised spellings</strong> (F3breze, F.e.b.r.e.z.e, lookalike letters); turn on <strong>fuzzy matching</strong> to also allow a few wrong letters, and use the strictness slider to choose how many. Matches found this way are tagged Disguised or Fuzzy in the "Matched as" column.
      Set a <strong>counterfeit review threshold</strong> to also flag sellers whose reviews frequently mention fake/counterfeit/knockoff products &mdash; this works across China, US, and Unknown alike, as long as the keyword still matches.
      Switch <strong>Search by</strong> to "Listing Brand" to instead pull every listing tagged with a specific seller-supplied Brand attribute (e.g. a private bulk-seller brand code) &mdash; brand mode trusts the listing's own Brand field exactly, it does not require the brand text to appear in the title.
      Set <strong>Target country</strong> to "Any (no country filter)" to pull every listing matching your keyword/brand regardless of country signal &mdash; useful for a pure brand sweep where you just want everything under that brand, not just the China/US-flagged subset.
      Your <strong>Market</strong>, <strong>Target country</strong>, <strong>Search by</strong>, and <strong>Listing brand matching</strong> choices are remembered in this browser between visits.
    </p>
  </div>
"""
    return html_page("New scan", body, SEARCH_FORM_SCRIPT, active="scan", scripts=("match_options.js",),
                     subtitle="Best-effort listing prescreening for seller-country signals. If the country is not explicit, the tool says unknown.")


@app.get("/scan", response_class=HTMLResponse)
def scan(
    keywords: str = Query(default=""),
    country: str = Query(default=DEFAULT_COUNTRY),
    market: str = Query(default=DEFAULT_MARKET),
    search_type: str = Query(default=DEFAULT_SEARCH_TYPE),
    limit: int = Query(default=DEFAULT_LIMIT, ge=0, le=500),
    workers: int = Query(default=LISTING_WORKERS, ge=1, le=MAX_LISTING_WORKERS),
    mode: str = Query(default=DEFAULT_MODE),
    max_price: str = Query(default=""),
    price_only: bool = Query(default=False),
    match_description: bool = Query(default=False),
    counterfeit_min: str = Query(default=""),
    negative_keywords: str = Query(default=""),
    generic_mode: Literal["standard", "only", "include"] = Query(default="standard"),
    generic_only: bool = Query(default=False),  # legacy bookmarked scan URLs
    fuzzy: bool = Query(default=False),
    fuzzy_strictness: int = Query(default=DEFAULT_STRICTNESS, ge=0, le=100),
    disguised_search: bool = Query(default=False),
) -> HTMLResponse:
    terms = split_keywords(keywords)
    if not terms:
        return home()
    parsed_max_price: float | None = None
    if max_price.strip():
        try:
            parsed_max_price = max(0.0, float(max_price.strip()))
        except ValueError:
            parsed_max_price = None
    parsed_counterfeit_min: int | None = None
    if counterfeit_min.strip():
        try:
            parsed_counterfeit_min = max(0, int(float(counterfeit_min.strip())))
        except ValueError:
            parsed_counterfeit_min = None
    selected_generic_mode = "only" if generic_only else generic_mode
    job_id = create_job(terms, country, limit, workers, mode, parsed_max_price, price_only, match_description, parsed_counterfeit_min, market, search_type, split_keywords(negative_keywords), selected_generic_mode == "only",
                        fuzzy_strictness if fuzzy else None, disguised_search, selected_generic_mode == "include")
    threading.Thread(target=run_job, args=(job_id,), daemon=True).start()
    return RedirectResponse(url=f"/job/{job_id}", status_code=302)


@app.get("/job/{job_id}", response_class=HTMLResponse)
def job_page(job_id: str) -> HTMLResponse:
    if job_id not in JOBS:
        summary = scan_history.scan_summary(DB_PATH, job_id)
        if not summary:
            return HTMLResponse("Scan not found", status_code=404)
        safe_id = html_escape(job_id, quote=True)
        state = summary["status"] or "Finished or stopped before restart (status not recorded)"
        body = f'''<div class="card">
          <p>This scan is archived. Live progress and logs were kept only while the server was running;
             restarting it does not delete saved results or review decisions.</p>
          <div class="summary"><span class="pill">Scan: {safe_id}</span>
            <span class="pill">Started: {html_escape(summary["created_at"])}</span>
            <span class="pill">Status: {html_escape(state)}</span>
            <span class="pill">Saved rows: {summary["rows"]:,}</span>
            <span class="pill">Confirmed matches: {summary["matches"]:,}</span></div>
          <div class="row-actions"><a class="button" href="/queue?job={quote_plus(job_id)}&amp;country_signal=all">View saved results</a>
            <a class="button secondary" href="/history">Scan history</a>
            <a class="button secondary" href="/">New scan</a></div>
        </div>'''
        return html_page("Live scan", body, active="job", subtitle="Saved scan summary; live updates resume with a new scan.")
    body = f"""
  <div class="card">
    <div class="summary">
      <span class="pill" id="job-id">Job: {html_escape(job_id)}</span>
      <span class="pill" id="job-market">Market: loading...</span>
      <span class="pill" id="job-status">Status: loading...</span>
      <span class="pill" id="job-message">Message: starting...</span>
      <span class="pill" id="china-count">China signal: 0</span>
      <span class="pill" id="us-count">US / explicit non-China signal: 0</span>
      <span class="pill" id="unknown-count">Unknown signal: 0</span>
      <span class="pill" id="excluded-count" hidden>Excluded by title: 0</span>
      <span class="pill" id="generic-only-pill" hidden title="Matches need a placeholder Brand (Generic, Unbranded, NOBRAND, Unknown, ...). Any country of origin counts.">Generic brands only &middot; any country</span>
      <span class="pill" id="generic-include-pill" hidden title="Regular listings need the target country; generic listing brands count from any country.">Standard + generic &middot; any country for generic</span>
      <span class="pill" id="fuzzy-pill" hidden>Fuzzy matching</span>
      <span class="pill" id="disguised-search-pill" hidden>Disguised spellings searched</span>
    </div>
    <div class="status">
      <div><strong>What the app is doing right now</strong></div>
      <div id="job-current" class="tiny">Waiting for job data...</div>
      <div class="bar" id="progress-bar-wrap"><div id="progress-bar"></div></div>
      <div id="job-stats" class="tiny"></div>
    </div>
    <div class="top-actions">
      <form id="stop-form" action="/job/{html_escape(job_id)}/stop" method="post" style="display:inline; margin:0;">
        <button type="submit" class="button secondary">Stop Scan</button>
      </form>
      <form id="resume-form" action="/job/{html_escape(job_id)}/resume" method="post" style="display:none; margin:0;">
        <button type="submit" class="button">Resume Scan</button>
      </form>
      <a class="button" href="/queue?job={html_escape(job_id)}{'&amp;country_signal=all' if JOBS.get(job_id, {}).get('generic_only') or JOBS.get(job_id, {}).get('generic_include') else ''}">Open review queue</a>
      <a class="button secondary" href="/queue?job={html_escape(job_id)}">Choose export</a>
      <a class="button secondary" href="/">New scan</a>
    </div>
    <div class="card inset" style="margin:0;">
      <div><strong>Live log</strong></div>
      <div id="job-logs" class="logbox"></div>
    </div>
    <div class="card" style="margin-top:16px;">
      <div><strong>Latest scan rows</strong> <span class="tiny">Most recent scan results, split by country signal</span></div>
      <label style="display:flex; align-items:center; gap:8px; margin-top:8px; font-weight:600;">
        <input id="matches-only-toggle" type="checkbox" checked style="width:auto;">
        Show confirmed matches only <span class="tiny" style="font-weight:400;">(uncheck to see every scanned listing, matched or not)</span>
      </label>
      <div class="tiny" id="latest-china-count" style="margin-top:8px;">China: 0</div>
      <div class="table-scroll">
        <table id="latest-china-table"></table>
      </div>
      <details style="margin-top:14px;">
        <summary class="tiny" id="latest-us-count">US signal: 0</summary>
        <div class="table-scroll" style="margin-top:8px;">
          <table id="latest-us-table"></table>
        </div>
      </details>
      <details style="margin-top:10px;">
        <summary class="tiny" id="latest-unknown-count">Unknown signal: 0</summary>
        <div class="table-scroll" style="margin-top:8px;">
          <table id="latest-unknown-table"></table>
        </div>
      </details>
    </div>
    <div class="card" style="margin-top:16px;">
      <div class="toolbar">
        <div class="tiny">All scanned listings are split into China, US, and Unknown sections. Unknown should mean genuinely unclear.</div>
      </div>
      <div class="card inset" style="margin: 0 0 12px 0; padding: 12px;">
        <div class="grid" style="grid-template-columns: 1.3fr .8fr .8fr; align-items:end;">
          <div>
            <label for="filter-text">Filter text</label>
            <input id="filter-text" type="text" placeholder="Keyword, seller, listing, evidence...">
          </div>
          <div>
            <label for="filter-status">Status</label>
            <select id="filter-status">
              <option value="all">All</option>
              <option value="match">Match</option>
              <option value="unknown">Unknown</option>
              <option value="error">Error</option>
              <option value="blocked">Blocked</option>
              <option value="excluded">Excluded</option>
            </select>
          </div>
          <div>
            <button type="button" class="button secondary" id="clear-filters">Clear Filters</button>
          </div>
        </div>
      </div>
      <h3 style="margin: 12px 0 6px; display:flex; justify-content:space-between; align-items:center; gap:12px;">
        <span>China signal (all scanned, check the Status column for confirmed matches)</span>
        <button type="button" class="button secondary" id="toggle-china">Hide Results</button>
      </h3>
      <table id="china-table" data-resizable>
        <thead>
          <tr>
            <th>Keyword</th><th>Listing</th><th>Seller Page</th><th>Seller</th><th>Ship From</th><th>Country</th><th>Status</th><th>Evidence</th>
          </tr>
        </thead>
        <tbody id="china-results-body">
          <tr><td colspan="8" class="muted">No China results yet. The scan is warming up.</td></tr>
        </tbody>
      </table>
      <h3 style="margin: 18px 0 6px; display:flex; justify-content:space-between; align-items:center; gap:12px;">
        <span>US / explicit non-China signal</span>
        <button type="button" class="button secondary" id="toggle-us">Show Results</button>
      </h3>
      <table id="us-table" data-resizable style="display:none;">
        <thead>
          <tr>
            <th>Keyword</th><th>Listing</th><th>Seller Page</th><th>Seller</th><th>Ship From</th><th>Country</th><th>Status</th><th>Evidence</th>
          </tr>
        </thead>
        <tbody id="us-results-body">
          <tr><td colspan="8" class="muted">No US results yet. The scan is warming up.</td></tr>
        </tbody>
      </table>
      <h3 style="margin: 18px 0 6px; display:flex; justify-content:space-between; align-items:center; gap:12px;">
        <span>Unknown country signal</span>
        <button type="button" class="button secondary" id="toggle-unknown">Show Results</button>
      </h3>
      <table id="unknown-table" data-resizable style="display:none;">
        <thead>
          <tr>
            <th>Keyword</th><th>Listing</th><th>Seller Page</th><th>Seller</th><th>Ship From</th><th>Country</th><th>Status</th><th>Evidence</th>
          </tr>
        </thead>
        <tbody id="unknown-results-body">
          <tr><td colspan="8" class="muted">No Unknown results yet. The scan is warming up.</td></tr>
        </tbody>
      </table>
    </div>
  </div>
"""
    extra_script = """
const jobId = __JOB_ID__;
function escapeHtml(text) {
  return String(text).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;').replace(/'/g,'&#39;');
}
function renderGroupedRows(rows, sectionKey) {
  const sorted = [...rows].sort((a, b) => {
    const sellerA = (a.seller_name || 'Unknown').toLowerCase();
    const sellerB = (b.seller_name || 'Unknown').toLowerCase();
    if (sellerA !== sellerB) return sellerA.localeCompare(sellerB);
    const keywordA = (a.keyword || '').toLowerCase();
    const keywordB = (b.keyword || '').toLowerCase();
    if (keywordA !== keywordB) return keywordA.localeCompare(keywordB);
    return (a.title || a.listing_url || '').localeCompare(b.title || b.listing_url || '');
  });

  const sellerGroups = [];
  for (const row of sorted) {
    const sellerKey = row.seller_name || 'Unknown';
    const lastSeller = sellerGroups[sellerGroups.length - 1];
    if (!lastSeller || lastSeller.key !== sellerKey) {
      sellerGroups.push({ key: sellerKey, rows: [row], keywords: [] });
    } else {
      lastSeller.rows.push(row);
    }
  }

  return sellerGroups.map(sellerGroup => {
    const sellerKey = `${sectionKey}::seller::${sellerGroup.key}`;
    const sellerCollapsed = collapsedGroups.has(sellerKey);
    const sellerId = sellerGroup.rows[0] && sellerGroup.rows[0].seller_page_url
      ? `<a href="${escapeHtml(sellerGroup.rows[0].seller_page_url)}" target="_blank" rel="noopener noreferrer">${escapeHtml(sellerGroup.rows[0].seller_id || 'Unknown')}</a>`
      : escapeHtml((sellerGroup.rows[0] && sellerGroup.rows[0].seller_id) || 'Unknown');

    const keywordGroups = [];
    for (const row of sellerGroup.rows) {
      const keywordKey = row.keyword || 'Unknown';
      const lastKeyword = keywordGroups[keywordGroups.length - 1];
      if (!lastKeyword || lastKeyword.key !== keywordKey) {
        keywordGroups.push({ key: keywordKey, rows: [row] });
      } else {
        lastKeyword.rows.push(row);
      }
    }

    const keywordHtml = keywordGroups.map(keywordGroup => {
      const keywordKey = `${sellerKey}::keyword::${keywordGroup.key}`;
      const keywordCollapsed = sellerCollapsed || collapsedGroups.has(keywordKey);
      return `
    <tr class="keyword-row" data-group-key="${escapeHtml(keywordKey)}">
      <td colspan="8">
        <button type="button" class="button secondary group-toggle" data-group-key="${escapeHtml(keywordKey)}">${keywordCollapsed ? 'Expand' : 'Collapse'}</button>
        <strong>${escapeHtml(keywordGroup.key)}</strong> <span class="tiny">(${keywordGroup.rows.length} result${keywordGroup.rows.length === 1 ? '' : 's'})</span>
      </td>
    </tr>
    ${keywordGroup.rows.map(r => `
    <tr class="group-item" data-group-key="${escapeHtml(sellerKey)} ${escapeHtml(keywordKey)}"${keywordCollapsed ? ' style="display:none;"' : ''}>
      <td>${escapeHtml(r.keyword)}</td>
      <td><a href="${escapeHtml(r.listing_url)}" target="_blank" rel="noopener noreferrer">${escapeHtml(r.title || r.listing_url)}</a></td>
      <td>${escapeHtml(r.seller_page_url || 'Unknown')}</td>
      <td>${escapeHtml(r.seller_name || 'Unknown')}</td>
      <td>${escapeHtml(r.ship_from_country || 'Unknown')}</td>
      <td>${escapeHtml(r.country_signal || 'Unknown')}</td>
      <td><span class="badge ${escapeHtml(r.scan_status || 'unknown')}">${escapeHtml(r.scan_status || 'unknown')}</span></td>
      <td>${escapeHtml(r.evidence || '')}</td>
    </tr>`).join('')}`;
    }).join('');

    return `
    <tr class="group-row" data-group-key="${escapeHtml(sellerKey)}">
      <td colspan="8">
        <button type="button" class="button secondary group-toggle" data-group-key="${escapeHtml(sellerKey)}">${sellerCollapsed ? 'Expand' : 'Collapse'}</button>
        <strong>${escapeHtml(sellerGroup.key)}</strong> <span class="tiny">(${sellerGroup.rows.length} result${sellerGroup.rows.length === 1 ? '' : 's'}) Seller ID: ${sellerId}</span>
      </td>
    </tr>
    ${keywordHtml}
  `;
  }).join('');
}

const collapsedGroups = new Set(JSON.parse(localStorage.getItem('collapsedGroups') || '[]'));
function persistCollapsedGroups() {
  localStorage.setItem('collapsedGroups', JSON.stringify([...collapsedGroups]));
}
function wireGroupToggles(scope) {
  scope.querySelectorAll('.group-toggle').forEach(btn => {
    btn.onclick = () => {
      const key = btn.dataset.groupKey;
      if (!key) return;
      if (collapsedGroups.has(key)) collapsedGroups.delete(key);
      else collapsedGroups.add(key);
      persistCollapsedGroups();
      refreshJob();
    };
  });
}
const latestTables = {};
['china', 'us', 'unknown'].forEach(key => {
  latestTables[key] = ListingTable.create(document.getElementById(`latest-${key}-table`), {
    pageSize: 50, storageKey: `prescreener_latest_${key}`, emptyText: () => (document.getElementById('matches-only-toggle') || {}).checked === false ? 'No scan rows yet.' : 'No confirmed matches yet.',
  });
});
const matchesOnlyToggleEl = document.getElementById('matches-only-toggle');
if (matchesOnlyToggleEl) {
  matchesOnlyToggleEl.addEventListener('change', refreshJob);
}

async function refreshJob() {
  const res = await fetch(`/api/job/${jobId}`);
  const data = await res.json();
  const stats = data.stats || {};
  document.getElementById('job-status').textContent = 'Status: ' + data.status;
  const resumeForm = document.getElementById('resume-form');
  const stopForm = document.getElementById('stop-form');
  if (resumeForm && stopForm) {
    const resumable = data.status === 'stopped' || data.status === 'error';
    resumeForm.style.display = resumable ? 'inline' : 'none';
    stopForm.style.display = resumable ? 'none' : 'inline';
  }
  document.getElementById('job-market').textContent = 'Market: ' + (data.market || 'US');
  document.getElementById('job-message').textContent = 'Message: ' + data.message;
  document.getElementById('job-current').textContent = data.status === 'stopping' ? 'Stopping scan...' : data.message;
  document.getElementById('job-stats').textContent = `Progress: ${data.current} / ${data.total} | Scanned: ${stats.scanned_total || 0} | Matches queued: ${stats.results_total || 0}`;
  const logs = document.getElementById('job-logs');
  logs.innerHTML = data.logs.map(line => '<div>' + line.replace(/&/g,'&amp;').replace(/</g,'&lt;') + '</div>').join('');
  logs.scrollTop = logs.scrollHeight;
  const scanned = data.scanned || [];
  const matchesOnlyToggle = document.getElementById('matches-only-toggle');
  const matchesOnly = matchesOnlyToggle ? matchesOnlyToggle.checked : true;
  const chinaRows = scanned.filter(r => (r.country_signal || 'Unknown') === 'China');
  const usRows = scanned.filter(r => (r.country_signal || 'Unknown') === 'US');
  const unknownRows = scanned.filter(r => (r.country_signal || 'Unknown') === 'Unknown');
  document.getElementById('china-count').textContent = `China signal: ${stats.china_total || 0}`;
  document.getElementById('us-count').textContent = `US / explicit non-China signal: ${stats.us_total || 0}`;
  document.getElementById('unknown-count').textContent = `Unknown signal: ${stats.unknown_total || 0}`;
  const excludedPill = document.getElementById('excluded-count');
  excludedPill.hidden = !(data.negative_terms && data.negative_terms.length);
  excludedPill.textContent = `Excluded by title: ${stats.excluded_total || 0}`;
  excludedPill.title = 'Excluded terms: ' + (data.negative_terms || []).join(', ');
document.getElementById('generic-only-pill').hidden = !data.generic_only;
document.getElementById('generic-include-pill').hidden = !data.generic_include;
const fuzzyPill = document.getElementById('fuzzy-pill');
fuzzyPill.hidden = data.fuzzy_strictness == null;
fuzzyPill.textContent = `Fuzzy matching: strictness ${data.fuzzy_strictness}`;
const disguisedPill = document.getElementById('disguised-search-pill');
disguisedPill.hidden = !data.disguised_searches;
disguisedPill.textContent = `+${data.disguised_searches} disguised spelling searches`;
disguisedPill.title = (data.search_plan || []).filter(p => p[0] !== p[1]).map(p => p[1]).join(', ');
  const visibleFor = rows => matchesOnly ? rows.filter(r => (r.scan_status || 'unknown') === 'match') : rows;
  const chinaVisible = visibleFor(chinaRows), usVisible = visibleFor(usRows), unknownVisible = visibleFor(unknownRows);
  const countSuffix = matchesOnly ? ' matches' : ' scanned';
  document.getElementById('latest-china-count').textContent = `China: ${chinaVisible.length}${countSuffix}`;
  document.getElementById('latest-us-count').textContent = `US: ${usVisible.length}${countSuffix}`;
  document.getElementById('latest-unknown-count').textContent = `Unknown: ${unknownVisible.length}${countSuffix}`;
  // newest first, like before (rows arrive oldest-first)
  latestTables.china.setRows(chinaVisible.slice().reverse());
  latestTables.us.setRows(usVisible.slice().reverse());
  latestTables.unknown.setRows(unknownVisible.slice().reverse());
  // Update shared state and rerender via the filter table's rerender()
  // (wired ONCE below, outside this polling function) instead of
  // re-attaching input/change listeners on every 1.2s poll. Stacking a
  // fresh listener per poll never gets cleaned up and each one holds a
  // full copy of the scanned rows via closure -- that's what was slowly
  // leaking memory and crashing the tab on long-running scans.
  filterState.chinaRows = chinaRows;
  filterState.usRows = usRows;
  filterState.unknownRows = unknownRows;
  rerenderFilteredTables();
  wireGroupToggles(document.getElementById('china-results-body'));
  wireGroupToggles(document.getElementById('us-results-body'));
  wireGroupToggles(document.getElementById('unknown-results-body'));
  // 'stopping' is deliberately NOT terminal -- keep polling until run_job
  // flips it to 'stopped', otherwise the Resume button never appears.
  const terminalStatuses = new Set(['done', 'stopped', 'error']);
  if (!terminalStatuses.has(data.status)) {
    setTimeout(refreshJob, 1200);
  }
}

const filterState = { chinaRows: [], usRows: [], unknownRows: [] };
function applyFilters(rows) {
  const filterTextEl = document.getElementById('filter-text');
  const filterText = (filterTextEl ? filterTextEl.value : '').trim().toLowerCase();
  const filterStatusEl = document.getElementById('filter-status');
  const filterStatus = filterStatusEl ? filterStatusEl.value : 'all';
  return rows.filter(r => {
    const haystack = [r.keyword, r.title, r.listing_url, r.seller_page_url, r.seller_name, r.ship_from_country, r.country_signal, r.evidence, r.scan_status].join(' ').toLowerCase();
    if (filterText && !haystack.includes(filterText)) return false;
    if (filterStatus !== 'all' && (r.scan_status || 'unknown') !== filterStatus) return false;
    return true;
  });
}
function rerenderFilteredTables() {
  const chinaBody = document.getElementById('china-results-body');
  const usBody = document.getElementById('us-results-body');
  const unknownBody = document.getElementById('unknown-results-body');
  const filteredChina = applyFilters(filterState.chinaRows);
  const filteredUS = applyFilters(filterState.usRows);
  const filteredUnknown = applyFilters(filterState.unknownRows);
  chinaBody.innerHTML = filteredChina.length ? renderGroupedRows(filteredChina, 'china') : '<tr><td colspan="8" class="muted">No China results match the filters.</td></tr>';
  usBody.innerHTML = filteredUS.length ? renderGroupedRows(filteredUS, 'us') : '<tr><td colspan="8" class="muted">No US results match the filters.</td></tr>';
  unknownBody.innerHTML = filteredUnknown.length ? renderGroupedRows(filteredUnknown, 'unknown') : '<tr><td colspan="8" class="muted">No Unknown results match the filters.</td></tr>';
  wireGroupToggles(chinaBody);
  wireGroupToggles(usBody);
  wireGroupToggles(unknownBody);
}
// Wired ONCE here -- not inside refreshJob -- so repeated polling never
// stacks duplicate listeners.
document.getElementById('toggle-us').onclick = () => {
  const usTable = document.getElementById('us-table');
  const toggleUS = document.getElementById('toggle-us');
  const hidden = usTable.style.display === 'none';
  usTable.style.display = hidden ? '' : 'none';
  toggleUS.textContent = hidden ? 'Hide Results' : 'Show Results';
};
document.getElementById('toggle-unknown').onclick = () => {
  const unknownTable = document.getElementById('unknown-table');
  const toggleUnknown = document.getElementById('toggle-unknown');
  const hidden = unknownTable.style.display === 'none';
  unknownTable.style.display = hidden ? '' : 'none';
  toggleUnknown.textContent = hidden ? 'Hide Results' : 'Show Results';
};
document.getElementById('clear-filters').onclick = () => {
  document.getElementById('filter-text').value = '';
  document.getElementById('filter-status').value = 'all';
  rerenderFilteredTables();
};
document.getElementById('filter-text').addEventListener('input', rerenderFilteredTables);
document.getElementById('filter-status').addEventListener('change', rerenderFilteredTables);

refreshJob();
initResizableTables();
initThumbPreview();
"""
    extra_script = RESIZABLE_TABLE_SCRIPT + THUMB_PREVIEW_SCRIPT + extra_script
    extra_script = extra_script.replace('__JOB_ID__', json.dumps(job_id))
    return html_page("Live scan", body, extra_script, active="job", scripts=("listing_table.js",),
                     subtitle="This page updates itself while the job runs.")




@app.post("/job/{job_id}/stop")
def stop_job(job_id: str) -> RedirectResponse:
    with lock:
        job = JOBS.get(job_id)
        # Only meaningful if a run_job thread is actually still alive to act on
        # the flag. If the job already reached a terminal state (done/error/
        # stopped) -- e.g. it finished a beat before the click landed -- there's
        # no thread left to ever flip "stopping" onward, which used to leave
        # the job stuck forever with no Resume button. Just no-op instead.
        if job and job.get("status") == "running":
            job["stop_requested"] = True
            job["status"] = "stopping"
            job["message"] = "Stopping scan..."
            append_log(job, "Stop requested by user.")
    return RedirectResponse(url=f"/job/{job_id}", status_code=302)


@app.post("/job/{job_id}/resume")
def resume_job(job_id: str) -> RedirectResponse:
    """Restart a stopped/errored job's run_job thread.

    Relies on the per-keyword `keyword_progress` + `seen_urls` state that
    run_job persists on every stop -- finished keywords are skipped
    entirely, unfinished ones pick up from their saved next search page
    instead of starting over from page 1.
    """
    with lock:
        job = JOBS.get(job_id)
        if not job or job.get("status") not in ("stopped", "error"):
            return RedirectResponse(url=f"/job/{job_id}", status_code=302)
        job["stop_requested"] = False
        job["status"] = "running"
        job["message"] = "Resuming scan..."
        job["finished_at"] = None
        append_log(job, "Resume requested by user -- picking up unfinished keywords where they left off.")
    threading.Thread(target=run_job, args=(job_id,), daemon=True).start()
    return RedirectResponse(url=f"/job/{job_id}", status_code=302)


@app.get("/api/match-preview")
def api_match_preview(keywords: str = Query(default=""), fuzzy: bool = Query(default=False),
                      strictness: int = Query(default=DEFAULT_STRICTNESS, ge=0, le=100)) -> JSONResponse:
    """Plain-English preview of how each typed keyword will be matched.

    Served from fuzzy_match.py so the New scan page never re-implements the rules.
    """
    terms = split_keywords(keywords)[:10]
    return JSONResponse({
        "keywords": [{"keyword": t, "words": [{"word": w, "allowed": allowed_edits(w, strictness if fuzzy else None), "short": len(w) <= SHORT_WORD_MAX}
                                              for w in keyword_words(t)]} for t in terms],
        "spellings": {t: disguised_spellings(t) for t in terms},
    })


@app.get("/api/job/{job_id}")
def api_job(job_id: str) -> JSONResponse:
    with lock:
        job = JOBS.get(job_id)
        if not job:
            return JSONResponse({"error": "job not found"}, status_code=404)
        payload = {k: v for k, v in job.items() if k not in ("keywords", "seen_urls")}
    return JSONResponse(payload)


QUEUE_ROW_FIELDS = (
    "id", "item_id", "image_url", "price_text", "price_value", "keyword", "title", "listing_url",
    "seller_id", "seller_name", "seller_page_url", "country_signal", "counterfeit_review_count",
    "scan_status", "evidence", "review_state", "match_kind", "remembered",
)


def _queue_scope(job: str) -> str | None:
    if job != "latest":
        return None if job == "all" else job
    if LATEST_JOB_ID:
        return LATEST_JOB_ID
    # Live job state is in memory; after a restart, use the last saved scan.
    with sqlite3.connect(DB_PATH) as conn:
        row = conn.execute("select job_id from reviews order by id desc limit 1").fetchone()
    return row[0] if row else "__no_scan__"


@app.get("/api/queue")
def api_queue(state: str = Query(default="pending"), country_signal: str = Query(default="China"), job: str = Query(default="latest")) -> JSONResponse:
    """Rows for the queue table. Only the fields the table shows, to keep big views light."""
    rows = get_reviews(state, country_signal, job_id=_queue_scope(job))
    return JSONResponse([{k: row.get(k) for k in QUEUE_ROW_FIELDS} for row in rows])


QUEUE_TABLE_SCRIPT = """
FastReview.init({ params: __QUEUE_PARAMS__ });
initThumbPreview();
"""

BULK_ACTIONS = {"approve": "approved", "reject": "rejected", "unset": "pending"}


@app.post("/api/review/bulk")
def api_review_bulk(action: str = Body(...), ids: list[int] = Body(...)) -> JSONResponse:
    """Fast Review: approve / dispute (reject) / unset many queue rows in one call."""
    if action not in BULK_ACTIONS:
        return JSONResponse({"error": f"unknown action {action!r}"}, status_code=400)
    ids = list(dict.fromkeys(ids))[:5000]
    updated, cascaded = update_review_states(ids, BULK_ACTIONS[action])
    return JSONResponse({"updated": updated, "cascaded": cascaded, "review_state": BULK_ACTIONS[action]})


def export_controls(job: str, state: str = "all") -> str:
    options = "".join(
        f'<option value="{key}"{" selected" if key == state else ""}>{label}</option>'
        for key, label in EXPORT_SCOPES.items()
    )
    return (f'<form action="/export.xlsx" method="get" class="export-controls">'
            f'<input type="hidden" name="job" value="{html_escape(job, quote=True)}">'
            f'<label for="export-scope">Export</label>'
            f'<select id="export-scope" name="scope">{options}</select>'
            f'<button type="submit" class="button secondary">Download XLSX</button>'
            f'<button type="button" id="export-selected" class="button secondary" disabled>Export selected (0)</button>'
            f'</form>')


@app.get("/history", response_class=HTMLResponse)
def history_page() -> HTMLResponse:
    return scan_history_page.render_history(DB_PATH, html_page)


@app.get("/queue", response_class=HTMLResponse)
def queue_page(state: str = Query(default="pending"), country_signal: str = Query(default="China"), job: str = Query(default="latest")) -> HTMLResponse:
    job_filter = _queue_scope(job)
    scope_note = ("Showing results from your most recent scan only." if job == "latest" else
                  "Showing results from ALL scans ever run (full history)." if job == "all" else
                  "Showing results from this scan only.")
    scope_toggle = (
        f'<a class="button secondary" href="/queue?state={state}&country_signal={country_signal}&job=all">View all scans (history)</a>'
        if job == "latest" else
        f'<a class="button secondary" href="/queue?state={state}&country_signal={country_signal}&job=latest">Back to latest scan only</a>'
    )
    state_counts = {s: count_reviews(s, country_signal, job_id=job_filter) for s in ["pending", "approved", "rejected", "not_queued"]}
    country_counts = {c: count_reviews(state, c, job_id=job_filter) for c in ["China", "US", "Unknown", "all"]}

    def _filter_link(href: str, label: str, is_current: bool) -> str:
        current = ' active" aria-current="true' if is_current else ""
        return f'<a class="button secondary{current}" href="{href}">{label}</a>'

    state_links = "".join(
        _filter_link(f"/queue?state={s}&country_signal={country_signal}&job={job}", f"{STATE_TAB_LABELS[s]} ({state_counts[s]})", s == state)
        for s in state_counts
    )
    country_links = "".join(
        _filter_link(f"/queue?state={state}&country_signal={c}&job={job}", f"{c if c.isupper() else c.title()} ({country_counts[c]})", c == country_signal)
        for c in ["China", "US", "Unknown", "all"]
    )
    body = f"""
  <div class="card">
    <div class="top-actions">
      <a class="button" href="/">New scan</a>
      {export_controls(job, state if state in EXPORT_SCOPES else 'all')}
      <span class="tiny">Exports use the scan shown and your chosen category; selected export uses checked rows within that category. Country and column filters do not apply to whole-category exports.</span>
      {scope_toggle}
    </div>
    <p class="tiny" style="margin: 6px 0 0;">{html_escape(scope_note)}</p>
    <div class="row-actions">
      <div class="tiny" style="width:100%; margin-bottom:6px;">Review status</div>
      {state_links}
    </div>
    <div class="row-actions">
      <div class="tiny" style="width:100%; margin-bottom:6px;">Country view</div>
      {country_links}
    </div>
    <p class="tiny" style="margin-top:12px;">Click a column name to sort. Use the boxes under the headers to filter. Hover a status badge to see its evidence.
      Decisions are remembered by item ID: dispute an item once and every repeat of it (now and in future scans) is auto-disputed; approved items come back marked "approved before". <strong>Not set</strong> forgets the decision.</p>
    <div class="fr-panel">
      <div class="fr-toolbar" role="toolbar" aria-label="Fast review">
        <div class="fr-count"><span id="fr-rowcount">Loading...</span> <span id="fr-selected" class="fr-selected" hidden></span></div>
        <div class="fr-actions">
          <button type="button" class="fr-icon fr-approve" data-bulk="approve" title="Approve selected (Shift+A)" aria-label="Approve selected" aria-keyshortcuts="Shift+A" disabled>
            <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M5 13l4 4L19 7"/></svg><span>Approve</span></button>
          <button type="button" class="fr-icon fr-dispute" data-bulk="reject" title="Dispute selected (Shift+D)" aria-label="Dispute selected" aria-keyshortcuts="Shift+D" disabled>
            <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M6 18L18 6M6 6l12 12"/></svg><span>Dispute</span></button>
          <button type="button" class="fr-icon fr-unset" data-bulk="unset" title="Not set: send back to pending (Shift+Z)" aria-label="Not set, send back to pending" aria-keyshortcuts="Shift+Z" disabled>
            <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M18.364 18.364A9 9 0 005.636 5.636m12.728 12.728A9 9 0 015.636 5.636m12.728 12.728L5.636 5.636"/></svg><span>Not set</span></button>
          <button type="button" class="fr-icon" id="fr-clear" title="Clear selection (Esc)" aria-label="Clear selection" aria-keyshortcuts="Escape" disabled>
            <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M4 7h16M10 11v6M14 11v6M5 7l1 12a2 2 0 002 2h8a2 2 0 002-2l1-12M9 7V4h6v3"/></svg><span>Clear</span></button>
          <button type="button" class="fr-icon" id="fr-help-toggle" aria-expanded="false" aria-controls="fr-help" title="Keyboard shortcuts">
            <svg viewBox="0 0 24 24" aria-hidden="true"><path d="M3 7h18v10H3zM7 11h.01M11 11h.01M15 11h.01M8 14h8"/></svg><span>Keys</span></button>
        </div>
      </div>
      <div id="fr-help" class="fr-help tiny" hidden>
        <strong>Click</strong> a box to select &middot; <strong>Shift+click</strong> another box to select the whole range (Shift+click on a ticked range unticks it) &middot;
        <kbd>J</kbd>/<kbd>K</kbd> move down/up &middot; <kbd>Space</kbd> toggle row &middot; <kbd>Shift</kbd>+<kbd>J</kbd>/<kbd>K</kbd> extend selection &middot;
        <kbd>Shift</kbd>+<kbd>A</kbd> approve &middot; <kbd>Shift</kbd>+<kbd>D</kbd> dispute &middot; <kbd>Shift</kbd>+<kbd>Z</kbd> not set &middot; <kbd>Esc</kbd> clear.
        The header box selects every row that matches your filters, not just the ones on screen.
      </div>
      <div id="fr-status" class="fr-status" role="status" aria-live="polite"></div>
      <div class="fr-scroll" tabindex="0" aria-label="Review queue table, scrollable"><table id="queue-table"></table></div>
    </div>
  </div>
"""
    params = json.dumps({"state": state, "country_signal": country_signal, "job": job}).replace("</", "<\\/")
    script = THUMB_PREVIEW_SCRIPT + QUEUE_TABLE_SCRIPT.replace("__QUEUE_PARAMS__", params)
    return html_page("Review queue", body, script, active="queue", scripts=("listing_table.js", "fast_review.js"),
                     subtitle="Approve or dispute prescreened listings, then export the evidence.")



@app.get("/review/{record_id}/{action}")
def review_action(record_id: int, action: str, return_to: str = "/queue") -> RedirectResponse:
    if action not in {"approve", "reject"}:
        return RedirectResponse(url=return_to, status_code=302)
    update_review_state(record_id, "approved" if action == "approve" else "rejected")
    return RedirectResponse(url=return_to, status_code=302)


@app.get("/export.xlsx")
def export_xlsx(job: str = Query(default="latest"), scope: ExportScope = Query(default="all")) -> FileResponse:
    path = export_reviews_xlsx(job_id=_queue_scope(job), scope=scope)
    return _xlsx_response(path, scope)


@app.post("/export.xlsx")
def export_selected_xlsx(ids: list[int] = Body(...), job: str = Query(default="latest"),
                         scope: ExportScope = Query(default="all")) -> FileResponse:
    """Export checked queue records only, within the requested scan and scope."""
    if len(ids) > 5000:
        return JSONResponse({"error": "Select at most 5,000 rows per export"}, status_code=400)
    path = export_reviews_xlsx(job_id=_queue_scope(job), scope=scope, selected_ids=set(ids))
    return _xlsx_response(path, scope)


def _xlsx_response(path: Path, scope: ExportScope) -> FileResponse:
    return FileResponse(
        path,
        filename=f"prescreener_{scope}.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        background=BackgroundTask(path.unlink, missing_ok=True),
    )


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="127.0.0.1", port=8001, reload=False)
