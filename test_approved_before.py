"""Approved Before bucket + category inheritance, on a throwaway database."""
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import decision_memory
import marketplace_prescreener as app

CAT = "IP Rights Infringement Power Tools"
OTHER = "IP Rights Infringement Toys"


class ApprovedBeforeTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        patcher = patch.object(app, "DB_PATH", Path(folder.name) / "t.sqlite3")
        patcher.start()
        self.addCleanup(patcher.stop)
        app.init_db()

    def add(self, rid, item, state="pending", category="", remembered="", job="scan-a"):
        with sqlite3.connect(app.DB_PATH) as conn:
            conn.execute("""insert into reviews (id, job_id, item_id, scan_status, review_state, review_category,
                remembered, keyword, search_url, listing_url, title, seller_name, ship_from_country,
                target_country, country_signal, evidence, created_at)
                values (?, ?, ?, 'match', ?, ?, ?, 'drill', '', '', 'drill', 'S', '', 'China', 'China', '', '')""",
                         (rid, job, item, state, category, remembered))

    def row(self, rid):
        with sqlite3.connect(app.DB_PATH) as conn:
            return conn.execute("select review_state, review_category, remembered from reviews where id = ?",
                                (rid,)).fetchone()

    def record(self, item):
        return app.ListingRecord(job_id="scan-new", keyword="drill", search_url="", listing_url="", seller_page_url="",
                                 title="drill", seller_name="S", seller_id="", ship_from_country="", target_country="China",
                                 country_signal="China", evidence_source="", evidence="", scan_status="match",
                                 item_id=item)

    def test_new_scan_repeat_goes_to_bucket_with_inherited_category(self):
        self.add(1, "111")
        app.update_review_states([1], "approved")
        app.review_categories.assign(app.DB_PATH, [1], CAT)
        rid = app.save_record(self.record("111"))
        self.assertEqual(self.row(rid), ("repeat", CAT, "approved_before"))
        fresh = app.save_record(self.record("999"))           # never decided: untouched
        self.assertEqual(self.row(fresh), ("pending", "", ""))

    def test_approving_cascades_existing_pending_copies_but_keeps_their_own_category(self):
        self.add(1, "111", category=CAT)
        self.add(2, "222", job="scan-b")
        self.add(3, "111", job="scan-b")
        self.add(4, "222", job="scan-c", category=OTHER)
        app.update_review_states([1, 2], "approved")
        self.assertEqual(self.row(3), ("repeat", CAT, "approved_before"))
        self.assertEqual(self.row(4), ("repeat", OTHER, "approved_before"))   # never overwritten

    def test_only_newest_repeat_per_item_stays_visible_older_copies_hidden_not_deleted(self):
        self.add(1, "111", category=CAT)
        for rid, job in ((2, "scan-b"), (3, "scan-c"), (4, "scan-d")):
            self.add(rid, "111", job=job)
        app.update_review_states([1], "approved")
        self.assertEqual([self.row(r)[0] for r in (2, 3, 4)], ["superseded", "superseded", "repeat"])
        newer = app.save_record(self.record("111"))         # a later scan finds it again
        self.assertEqual(self.row(newer)[0], "repeat")
        self.assertEqual(self.row(4)[0], "superseded")
        self.assertEqual(app.count_reviews("repeat"), 1)
        self.assertEqual(app.count_reviews(), 5)               # 4 seeded + 1 new: nothing deleted
        path = app.export_reviews_xlsx(job_id="scan-b")         # whole-scan export still has it
        self.addCleanup(path.unlink, missing_ok=True)
        from openpyxl import load_workbook
        sheet = list(load_workbook(path, read_only=True).active.values)
        self.assertEqual(sheet[1][sheet[0].index("Review Status")], "Approved Before (older copy)")
        app.update_review_states([1], "pending")                # Not set brings every copy back
        self.assertEqual({self.row(r)[0] for r in (2, 3, 4, newer)}, {"pending"})

    def test_tabs_and_history_show_items_and_rows_when_they_differ(self):
        for rid, item in ((1, "111"), (2, "111"), (3, "222")):
            self.add(rid, item, state="approved", job=f"scan-{rid}")
        page = app.queue_page(state="approved", country_signal="all", job="all").body.decode()
        self.assertIn("Approved (2 items · 3 rows)", page)
        self.assertIn("Pending (0)", page)
        self.assertEqual(app.count_label("X", 5, 5), "X (5)")
        self.assertEqual(app.count_label("X", 1378, 207), "X (207 items · 1,378 rows)")

    def test_not_set_returns_repeats_to_pending_and_dispute_auto_disputes_them(self):
        self.add(1, "111", category=CAT)
        self.add(2, "111", job="scan-b")
        app.update_review_states([1], "approved")
        app.update_review_states([1], "pending")              # Not set / Undo
        self.assertEqual(self.row(2), ("pending", CAT, ""))   # category kept: nothing wiped
        app.update_review_states([1], "approved")
        app.update_review_states([1], "rejected")
        self.assertEqual(self.row(2), ("rejected", "", "auto_disputed"))

    def test_repeat_rows_can_be_decided_categorised_and_exported(self):
        self.add(1, "111", category=CAT)
        self.add(2, "111", job="scan-b")
        app.update_review_states([1], "approved")
        self.assertEqual(app.review_categories.assign(app.DB_PATH, [2], OTHER), 1)
        self.assertEqual(app.count_reviews("repeat"), 1)
        path = app.export_reviews_xlsx(scope="repeat")
        self.addCleanup(path.unlink, missing_ok=True)
        from openpyxl import load_workbook
        sheet = list(load_workbook(path, read_only=True).active.values)
        status = sheet[0].index("Review Status")
        self.assertEqual([r[status] for r in sheet[1:]], ["Approved Before"])
        app.update_review_states([2], "approved")
        self.assertEqual(self.row(2)[0], "approved")

    def test_startup_migration_moves_old_pending_repeats_and_is_idempotent(self):
        self.add(1, "111", state="approved", category=CAT)
        self.add(2, "111", remembered="approved_before", job="scan-b")
        app.init_db()
        app.init_db()
        self.assertEqual(self.row(2), ("repeat", CAT, "approved_before"))

    def test_queue_and_history_show_the_new_bucket(self):
        self.add(1, "111", category=CAT)
        self.add(2, "111", job="scan-b")
        app.update_review_states([1], "approved")
        page = app.queue_page(state="repeat", country_signal="all", job="all").body.decode()
        self.assertIn("Approved Before (1)", page)
        self.assertIn("Approved Before (1)", app.history_page().body.decode())
        self.assertEqual(decision_memory.REPEAT_STATE, "repeat")


if __name__ == "__main__":
    unittest.main()
